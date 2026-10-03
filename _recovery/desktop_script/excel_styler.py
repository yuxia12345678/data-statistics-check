# -*- coding: utf-8 -*-
"""
Excel美化渲染模块 excel_styler
职责：
  1、根据配置里面的样式参数（颜色、边框、数字格式）构建openpyxl样式对象；
  2、实现表头样式、斑马纹交替底色、水平浅灰色边框；
  3、实现冻结首行；自适应中文列宽（区分全角中文字符与半角英文字符）；
  4、按字段分类设置对齐方式：文本左对齐；数值右对齐；日期、分类居中；
  样式全部从业务配置读取，不硬编码颜色与格式。
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.worksheet import Worksheet


def build_excel_style_from_config(style_cfg: dict):
    """
    根据配置字典构造openpyxl可复用样式对象
    :param style_cfg: json中style_setting节点字典
    :return: 样式对象字典
    """
    # 表头背景色、表头字体
    header_fill = PatternFill(start_color=style_cfg["header_bg_color"],
                              end_color=style_cfg["header_bg_color"], fill_type="solid")
    header_font = Font(bold=True, color=style_cfg["header_font_color"])
    # 奇偶行斑马纹底色
    odd_row_fill = PatternFill(start_color=style_cfg["odd_row_bg"], end_color=style_cfg["odd_row_bg"], fill_type="solid")
    even_row_fill = PatternFill(start_color=style_cfg["even_row_bg"], end_color=style_cfg["even_row_bg"], fill_type="solid")
    # 仅保留水平浅灰色边框线
    thin_h_line = Side(style="thin", color=style_cfg["horizontal_border_color"])
    h_border = Border(top=None, bottom=thin_h_line, left=None, right=None)

    return {
        "header_fill": header_fill,
        "header_font": header_font,
        "odd_fill": odd_row_fill,
        "even_fill": even_row_fill,
        "h_border": h_border
    }


def auto_fit_column_width(ws: Worksheet):
    """
    自适应计算工作表列宽；中文（全角）字符权重2，英文数字权重1；限制最大最小宽度
    :param ws: openpyxl工作表对象
    """
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                val = str(cell.value)
                # 遍历字符计算显示长度
                length = sum([2 if '\u4e00' <= c <= '\u9fff' else 1 for c in val])
                if length > max_len:
                    max_len = length
        # 最小宽度8；最大宽度45；额外预留3个字符边距
        final_w = max(8, min(max_len + 3, 45))
        ws.column_dimensions[col_letter].width = final_w


def apply_sheet_format(ws: Worksheet, style_cfg: dict,
                       num_cols: list, pct_cols: list, date_cols: list, cat_cols: list):
    """
    对单个工作表执行全套美化逻辑：冻结窗格、表头样式、斑马纹底色、边框、对齐、数字格式
    :param ws: worksheet对象
    :param style_cfg: 样式配置字典
    :param num_cols: 金额类字段名称列表（千分位小数格式）
    :param pct_cols: 百分比字段名称列表（百分比保留两位小数）
    :param date_cols: 日期字段列表（居中对齐）
    :param cat_cols: 分类字段列表（居中对齐）
    """
    # 读取样式对象
    style_dict = build_excel_style_from_config(style_cfg)
    header_fill = style_dict["header_fill"]
    header_font = style_dict["header_font"]
    odd_fill = style_dict["odd_fill"]
    even_fill = style_dict["even_fill"]
    h_border = style_dict["h_border"]

    header_row_idx = 1
    # 冻结首行，表头固定
    ws.freeze_panes = f"A{header_row_idx + 1}"
    # 获取表头行所有单元格的值，构建列名与列序号映射
    header_names = [cell.value for cell in ws[header_row_idx]]
    name2idx = {n: i+1 for i, n in enumerate(header_names) if n is not None}

    # 设置表头行样式：居中、加粗、表头底色
    for cell in ws[header_row_idx]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    max_row = ws.max_row
    # 遍历数据行设置斑马纹底色以及底部水平边框
    for r in range(header_row_idx + 1, max_row + 1):
        row_cells = ws[r]
        # 判断奇偶行选择底色
        row_fill = odd_fill if (r % 2 == 1) else even_fill
        for cell in row_cells:
            cell.fill = row_fill
            cell.border = h_border
        # 根据字段分类设置对齐方式和单元格数字格式
        for col_name in name2idx:
            c_idx = name2idx[col_name]
            cell = ws.cell(row=r, column=c_idx)
            if col_name in num_cols:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = style_cfg["amount_number_format"]
            elif col_name in pct_cols:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = style_cfg["percent_number_format"]
            elif col_name in date_cols or col_name in cat_cols:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
    # 执行自适应列宽
    auto_fit_column_width(ws)
