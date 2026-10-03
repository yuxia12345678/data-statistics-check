# -*- coding: utf-8 -*-
"""
交付前自检程序（独立复核，不复用 analyze.py 的任何计算结果）。

用法：
    python3 verify.py --file 结果.xlsx --source 附件1-合同开票及回款核对表.xlsx
    可选： --profile spec|example   --json report.json

设计原则：
  - 所有期望值都从"源文件"重新算一遍，再与结果文件的缓存值对照，
    形成源事实 → 产物 的独立验证链，不用生成逻辑自证。
  - 公式视图（data_only=False）核公式与样式；数值视图（data_only=True）核缓存值。
  - 任何一项不通过 → 退出码 1，并按题目"注意事项(4)"逐项列出。
"""

import argparse
import collections
import json
import re
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import config as C

TOL = 1e-6


# ================================================================ 源事实（独立重算）

def source_facts(path, profile=None):
    wb = load_workbook(path, data_only=True)
    ws = None
    for w in wb.worksheets:
        if w.sheet_state == "visible":
            ws = w
            break
    header_row = None
    for r in range(1, min(ws.max_row, 20) + 1):
        vals = [str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else ""
                for c in range(1, ws.max_column + 1)]
        if "合同号" in vals:
            header_row = r
            break
    col_of = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if v is not None and str(v).strip() and str(v).strip() not in col_of:
            col_of[str(v).strip()] = c

    raw = []          # [(row, {field: value})]
    for r in range(header_row + 1, ws.max_row + 1):
        rec = {}
        empty = True
        for f, c in col_of.items():
            v = ws.cell(r, c).value
            rec[f] = v
            if v not in (None, ""):
                empty = False
        if not empty:
            raw.append((r, rec))

    def num(v):
        return 0.0 if v in (None, "") else float(v)

    facts = {
        "header_row": header_row,
        "col_of": col_of,
        "rows": len(raw),
        "raw_last_row": raw[-1][0] if raw else header_row,
        "totals": {f: sum(num(rec[f]) for _r, rec in raw)
                   for f in C.MERGE_SUM_FIELDS if f in col_of},
        "raw_cells": {},
    }
    for r, rec in raw:
        for f, c in col_of.items():
            facts["raw_cells"][(r, c)] = (rec[f], ws.cell(r, c).number_format)

    def key(rec, f):
        v = rec.get(f)
        return "" if v is None else (v if isinstance(v, str) else str(v)).strip()

    # 合同号分组
    g = collections.OrderedDict()
    for r, rec in raw:
        code = key(rec, "合同号")
        d = g.setdefault(code, {"sums": collections.defaultdict(float), "rows": 0})
        d["rows"] += 1
        for f in C.MERGE_SUM_FIELDS:
            d["sums"][f] += num(rec.get(f))
    facts["contracts"] = g
    facts["contract_count"] = len(g)

    # 维度分组
    dims = {}
    for _s, dim in C.DIMENSIONS:
        b = collections.OrderedDict()
        for r, rec in raw:
            k = key(rec, dim)
            if not k:
                continue
            d = b.setdefault(k, {"sums": collections.defaultdict(float), "codes": set(), "rows": 0})
            d["rows"] += 1
            d["codes"].add(key(rec, "合同号"))
            for f in C.MERGE_SUM_FIELDS:
                d["sums"][f] += num(rec.get(f))
        dims[dim] = b
    facts["dims"] = dims

    # 未回款原因分类（非空去重）
    # example profile 会把空值填充为「未填写」，源事实必须按同一口径重算，
    # 否则 R5/R6 会拿 spec 口径去核 example 产物。
    fill_blank = bool(profile and profile.get("fill_blank_reason"))
    b = collections.OrderedDict()
    for r, rec in raw:
        k = key(rec, "未回款原因分类")
        if not k:
            if not fill_blank:
                continue
            k = "未填写"
        d = b.setdefault(k, {"unpaid": 0.0, "codes": set(), "rows": 0})
        d["unpaid"] += num(rec.get("开票未回款"))
        d["rows"] += 1
        d["codes"].add(key(rec, "合同号"))
    facts["reasons"] = b
    return facts


# ================================================================ 检查项

def close(a, b):
    if a is None or b is None:
        return a is None and b is None
    try:
        return abs(float(a) - float(b)) <= max(TOL, abs(float(b)) * 1e-9)
    except (TypeError, ValueError):
        return str(a) == str(b)


class Checker(object):
    def __init__(self, out_path, src_path, profile):
        self.profile = profile
        self.wbf = load_workbook(out_path, data_only=False)
        self.wbv = load_workbook(out_path, data_only=True)
        self.srcf = source_facts(src_path, profile)
        # 源文件存在带样式但无表头/无数据的"幽灵列"（本例 P..S），
        # 属于源文件自带状态；「原始数据」必须原样保留，故美化与核验都只覆盖有效列。
        self.raw_ncols = max(self.srcf["col_of"].values())
        self.src_wb = load_workbook(src_path, data_only=True)
        self.results = []

    def add(self, rid, desc, ok, detail=""):
        self.results.append({"id": rid, "desc": desc, "ok": bool(ok), "detail": detail})

    # ---------- R1 工作表
    def r1_sheets(self):
        want = self.profile["sheet_order"]
        got = self.wbf.sheetnames
        self.add("R1", "8 张工作表名称与顺序严格匹配", got == want,
                 "期望=%s 实际=%s" % (want, got))

    # ---------- R2 原始数据原样
    def r2_raw_intact(self):
        ws = self.wbv[C.SHEET_RAW]
        src_ws = None
        for w in self.src_wb.worksheets:
            if w.sheet_state == "visible":
                src_ws = w
                break
        diff_val, diff_nf, n = [], [], 0
        col_of = self.srcf["col_of"]
        for (r, c), (v, nf) in self.srcf["raw_cells"].items():
            n += 1
            gv = ws.cell(r, c).value
            gnf = ws.cell(r, c).number_format
            if isinstance(v, (int, float)) and isinstance(gv, (int, float)):
                if not close(v, gv):
                    diff_val.append((r, c, v, gv))
            elif (v if v is not None else "") != (gv if gv is not None else ""):
                diff_val.append((r, c, repr(v)[:40], repr(gv)[:40]))
            if nf != gnf:
                diff_nf.append((r, c, nf, gnf))
        self.add("R2a", "「原始数据」全部单元格值与源文件一致（含行列顺序）",
                 not diff_val, "比对 %d 格，差异 %d 处 %s" % (n, len(diff_val), diff_val[:5]))
        self.add("R2b", "「原始数据」数字格式未被修改", not diff_nf,
                 "差异 %d 处 %s" % (len(diff_nf), diff_nf[:5]))
        hr = self.srcf["header_row"]
        heads_ok = all(ws.cell(hr, c).value == src_ws.cell(hr, c).value
                       for c in range(1, max(col_of.values()) + 1))
        self.add("R2c", "「原始数据」表头文字与行位置未变", heads_ok,
                 "表头行=%d" % hr)
        self.add("R2d", "「原始数据」行数与源一致（无增删行）",
                 ws.max_row == src_ws.max_row,
                 "源=%d 结果=%d" % (src_ws.max_row, ws.max_row))

    # ---------- R3 禁止硬编码
    def r3_formulas(self):
        stat_sheets = [s for s in self.profile["sheet_order"] if s != C.SHEET_RAW]
        no_formula, no_cache, total = [], [], 0
        for name in stat_sheets:
            wsf, wsv = self.wbf[name], self.wbv[name]
            for r in range(3, wsf.max_row + 1):
                for c in range(1, wsf.max_column + 1):
                    cell = wsf.cell(r, c)
                    v = cell.value
                    if v is None or v == "":
                        continue
                    # 维度名/合同号/指标名等键值是标识符，允许为字面量
                    if isinstance(v, str) and not v.startswith("="):
                        continue
                    total += 1
                    if not (isinstance(v, str) and v.startswith("=")):
                        no_formula.append((name, cell.coordinate, repr(v)[:40]))
                    if wsv.cell(r, c).value is None:
                        no_cache.append((name, cell.coordinate))
        # 数值列必须是公式：逐列判定
        numeric_hardcoded = []
        for name in stat_sheets:
            wsf, wsv = self.wbf[name], self.wbv[name]
            for r in range(3, wsf.max_row + 1):
                for c in range(1, wsf.max_column + 1):
                    cached = wsv.cell(r, c).value
                    if isinstance(cached, (int, float)) and cached not in (None,):
                        f = wsf.cell(r, c).value
                        if not (isinstance(f, str) and f.startswith("=")):
                            numeric_hardcoded.append((name, wsf.cell(r, c).coordinate, cached))
        self.add("R3a", "全部数值单元格均为动态公式（禁止硬编码）",
                 not numeric_hardcoded, "硬编码 %d 处 %s" % (len(numeric_hardcoded),
                                                          numeric_hardcoded[:6]))
        self.add("R3b", "公式单元格均带缓存值（程序化读取不为空）", not no_cache,
                 "缺缓存 %d 处 %s" % (len(no_cache), no_cache[:6]))
        ok = self._all_formulas_reference_raw(stat_sheets)
        self.add("R3c", "公式全部动态引用「原始数据」区间（或引用同表公式格派生）",
                 ok, "存在未引用数据源的公式：%s" % (getattr(self, "_bad_formula", None),))

    # 不含数据源引用的公式，只允许两种形态：
    #   1) 行号序号          ROW()-2
    #   2) 同表派生比率       IF(C3=0,"",E3/C3)  ← 分子分母本身都是引用数据源的公式格
    LOCAL_OK = (re.compile(r"^ROW\(\)-\d+$"),
                re.compile(r'^IF\(\$?[A-Z]{1,3}\$?\d+=0,"",\$?[A-Z]{1,3}\$?\d+/\$?[A-Z]{1,3}\$?\d+\)$'))

    def _all_formulas_reference_raw(self, stat_sheets):
        for name in stat_sheets:
            wsf = self.wbf[name]
            for row in wsf.iter_rows(min_row=3):
                for cell in row:
                    v = cell.value
                    if not (isinstance(v, str) and v.startswith("=")):
                        continue
                    body = v[1:]
                    if "原始数据" in body:
                        continue
                    if not any(p.match(body) for p in self.LOCAL_OK):
                        self._bad_formula = (name, cell.coordinate, body[:60])
                        return False
        return True

    # ---------- R4 公式错误
    def r4_errors(self):
        found = []
        for name in self.wbv.sheetnames:
            ws = self.wbv[name]
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value in (
                            "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!"):
                        found.append((name, cell.coordinate, cell.value))
        self.add("R4", "无公式错误值", not found, "错误 %d 处 %s" % (len(found), found[:8]))

    # ---------- R5 行数
    def r5_rowcounts(self):
        f = self.srcf
        ws = self.wbv[C.SHEET_MERGE]
        self.add("R5a", "合并表行数 = 不同合同号数", ws.max_row - 2 == f["contract_count"],
                 "结果=%d 期望=%d" % (ws.max_row - 2, f["contract_count"]))
        for sheet, dim in C.DIMENSIONS:
            ws = self.wbv[sheet]
            self.add("R5-" + dim, "%s 行数 = %s 唯一值数" % (sheet, dim),
                     ws.max_row - 2 == len(f["dims"][dim]),
                     "结果=%d 期望=%d" % (ws.max_row - 2, len(f["dims"][dim])))
        ws = self.wbv[C.SHEET_OVERVIEW]
        self.add("R5-总览", "统计总览为 8 项全局指标", ws.max_row - 2 == 8,
                 "结果=%d" % (ws.max_row - 2))
        ws = self.wbv[C.SHEET_REASON]
        self.add("R5-原因", "未回款原因分类汇总表行数 = 非空去重分类数",
                 ws.max_row - 2 == len(f["reasons"]),
                 "结果=%d 期望=%d（源字段全为空时为 0）" % (ws.max_row - 2, len(f["reasons"])))

    # ---------- R6 数值正确
    def r6_values(self):
        f = self.srcf
        # 合并表：按合同号核对 5 个求和字段
        ws = self.wbv[C.SHEET_MERGE]
        headers = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        idx = {h: i + 1 for i, h in enumerate(headers)}
        bad = []
        for r in range(3, ws.max_row + 1):
            code = ws.cell(r, idx["合同号"]).value
            g = f["contracts"].get(str(code).strip())
            if g is None:
                bad.append((r, code, "源中不存在该合同号"))
                continue
            for fld in C.MERGE_SUM_FIELDS:
                if fld not in idx:
                    continue
                if not close(ws.cell(r, idx[fld]).value, g["sums"][fld]):
                    bad.append((r, code, fld, ws.cell(r, idx[fld]).value, g["sums"][fld]))
        self.add("R6a", "合并表 5 个数值字段求和正确（463 合同 × 5 字段）",
                 not bad, "差异 %d 处 %s" % (len(bad), bad[:5]))

        # 维度表 7 项指标
        for sheet, dim in C.DIMENSIONS:
            ws = self.wbv[sheet]
            headers = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
            bad = []
            total_amt = f["totals"]["合同金额"]
            for r in range(3, ws.max_row + 1):
                k = str(ws.cell(r, 2).value).strip()
                d = f["dims"][dim].get(k)
                if d is None:
                    bad.append((r, k, "源中无此维度值"))
                    continue
                exp = [d["sums"]["合同金额"], d["sums"]["开票金额"], d["sums"]["回款合计"],
                       d["sums"]["开票未回款"], len(d["codes"]),
                       (d["sums"]["回款合计"] / d["sums"]["合同金额"]) if d["sums"]["合同金额"] else None,
                       (d["sums"]["开票未回款"] / total_amt) if total_amt else None]
                got = [ws.cell(r, c).value for c in range(3, 10)]
                for i, (a, b) in enumerate(zip(got, exp)):
                    if not close(a, b):
                        bad.append((sheet, r, k, headers[i + 2], a, b))
            self.add("R6-" + dim, "%s 全部 7 项指标与源事实一致" % sheet, not bad,
                     "差异 %d 处 %s" % (len(bad), bad[:5]))

        # 统计总览
        ws = self.wbv[C.SHEET_OVERVIEW]
        exp = [f["totals"]["合同金额"], f["totals"]["开票金额"], f["totals"]["回款合计"],
               f["totals"]["开票未回款"],
               f["totals"]["回款合计"] / f["totals"]["合同金额"],
               f["contract_count"], f["rows"], len(f["reasons"])]
        got = [ws.cell(3 + i, 2).value for i in range(8)]
        bad = [(C.OVERVIEW_LABELS[i], got[i], exp[i]) for i in range(8) if not close(got[i], exp[i])]
        self.add("R6-总览", "统计总览 8 项指标与源事实一致", not bad, "差异 %s" % bad)

        # 原因表
        ws = self.wbv[C.SHEET_REASON]
        bad = []
        total_unpaid = f["totals"]["开票未回款"]
        for r in range(3, ws.max_row + 1):
            k = str(ws.cell(r, 2).value).strip()
            d = f["reasons"].get(k)
            if d is None:
                bad.append((r, k, "源中无此分类"))
                continue
            for col, e in ((3, d["unpaid"]), (4, len(d["codes"])), (5, d["rows"]),
                           (6, d["unpaid"] / total_unpaid if total_unpaid else None)):
                if not close(ws.cell(r, col).value, e):
                    bad.append((r, k, col, ws.cell(r, col).value, e))
        self.add("R6-原因", "未回款原因分类汇总 4 项指标与源事实一致", not bad,
                 "差异 %s（源字段全空时本项为空表，自动通过）" % bad[:5])

    # ---------- R7 排序
    def r7_order(self):
        for sheet, dim in C.DIMENSIONS:
            ws = self.wbv[sheet]
            vals = [ws.cell(r, 6).value or 0 for r in range(3, ws.max_row + 1)]
            self.add("R7-" + dim, "%s 按开票未回款降序" % sheet,
                     all(vals[i] >= vals[i + 1] for i in range(len(vals) - 1)),
                     "首尾=%s/%s" % (vals[:1], vals[-1:]))
        ws = self.wbv[C.SHEET_REASON]
        vals = [ws.cell(r, 3).value or 0 for r in range(3, ws.max_row + 1)]
        self.add("R7-原因", "未回款原因分类汇总按开票未回款降序",
                 all(vals[i] >= vals[i + 1] for i in range(len(vals) - 1)), "行数=%d" % len(vals))
        ws = self.wbv[C.SHEET_MERGE]
        heads = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        amt_col = heads.index("合同金额") + 1
        vals = [ws.cell(r, amt_col).value or 0 for r in range(3, ws.max_row + 1)]
        self.add("R7-合并", "合并表按合同金额降序（与附件2-示例一致）",
                 all(vals[i] >= vals[i + 1] - 1e-9 for i in range(len(vals) - 1)),
                 "行数=%d" % len(vals))
        # 序号连续
        for name in [C.SHEET_MERGE, C.SHEET_REASON] + [s for s, _d in C.DIMENSIONS]:
            ws = self.wbv[name]
            seq = [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)]
            self.add("R7seq-" + name, "%s 序号全新且连续（1..N）" % name,
                     seq == list(range(1, len(seq) + 1)), "前3=%s 末=%s" % (seq[:3], seq[-1:]))

    # ---------- R8 数字格式
    def r8_formats(self):
        money = self.profile["money_format"]
        bad = []
        for name in self.profile["sheet_order"]:
            if name == C.SHEET_RAW:
                continue
            ws = self.wbf[name]
            headers = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
            for c, h in enumerate(headers, start=1):
                kind = C.FIELD_KIND.get(str(h).replace("总", "") if h else "", None) or \
                       C.FIELD_KIND.get(h, "text")
                if name == C.SHEET_OVERVIEW and c == 2:
                    continue  # 逐行类型不同，单独核
                want = {"money": money, "pct": C.PCT_FORMAT, "int": C.INT_FORMAT,
                        "date": C.DATE_FORMAT_SRC}.get(kind)
                if not want:
                    continue
                for r in range(3, ws.max_row + 1):
                    nf = ws.cell(r, c).number_format
                    if nf != want:
                        bad.append((name, ws.cell(r, c).coordinate, nf, want))
                        break
        ws = self.wbf[C.SHEET_OVERVIEW]
        want_by_row = [money, money, money, money, C.PCT_FORMAT, C.INT_FORMAT,
                       C.INT_FORMAT, C.INT_FORMAT]
        for i, w in enumerate(want_by_row):
            nf = ws.cell(3 + i, 2).number_format
            if nf != w:
                bad.append((C.SHEET_OVERVIEW, ws.cell(3 + i, 2).coordinate, nf, w))
        dec = len(money.split(".")[1]) if "." in money else 0
        self.add("R8", "金额千分位+%d位小数、百分比2位小数、整数与日期格式合规"
                 % dec, not bad, "不合规 %s" % bad[:8])

    # ---------- R9 对齐
    def r9_align(self):
        align = self.profile["align_by_kind"]
        bad = []
        for name in self.profile["sheet_order"]:
            ws = self.wbf[name]
            hr = self.srcf["header_row"] if name == C.SHEET_RAW else 2
            nc = self.raw_ncols if name == C.SHEET_RAW else ws.max_column
            headers = [ws.cell(hr, c).value for c in range(1, nc + 1)]
            for c, h in enumerate(headers, start=1):
                kind = C.FIELD_KIND.get(h, "text")
                if name == C.SHEET_OVERVIEW:
                    kind = "text" if c == 1 else "num"
                want = align.get(kind, "left")
                for r in range(hr + 1, ws.max_row + 1):
                    got = ws.cell(r, c).alignment.horizontal
                    if got != want:
                        bad.append((name, ws.cell(r, c).coordinate, kind, got, want))
                        break
        self.add("R9", "文本左对齐 / 数值右对齐 / 日期与分类居中", not bad,
                 "不合规 %s" % bad[:8])

    # ---------- R10 表头样式
    def r10_header(self):
        bad = []
        for name in self.profile["sheet_order"]:
            ws = self.wbf[name]
            hr = self.srcf["header_row"] if name == C.SHEET_RAW else 2
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(hr, c)
                if cell.value in (None, ""):
                    continue
                fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
                colr = cell.font.color.rgb if cell.font and cell.font.color else None
                if fill != C.HEADER_FILL or not cell.font.b or colr != C.HEADER_FONT_COLOR \
                        or cell.alignment.horizontal != "center":
                    bad.append((name, cell.coordinate, fill, cell.font.b, colr,
                                cell.alignment.horizontal))
        self.add("R10", "表头低饱和度蓝底 + 白色加粗 + 水平居中", not bad,
                 "不合规 %s" % bad[:6])

    # ---------- R11 斑马纹
    def r11_zebra(self):
        bad = []
        for name in self.profile["sheet_order"]:
            ws = self.wbf[name]
            first = self.srcf["header_row"] + 1 if name == C.SHEET_RAW else 3
            for i, r in enumerate(range(first, ws.max_row + 1)):
                fill = ws.cell(r, 1).fill
                rgb = fill.fgColor.rgb if fill and fill.fill_type else None
                want = C.ZEBRA_FILL if i % 2 == 0 else None
                if rgb != want:
                    bad.append((name, r, rgb, want))
        self.add("R11", "数据区斑马纹交替行底色", not bad, "不合规 %s" % bad[:6])

    # ---------- R12 边框
    def r12_border(self):
        mode = self.profile["border_mode"]
        bad = []

        def st(side):
            return getattr(side, "style", None) if side is not None else None

        def colr(side):
            c = getattr(side, "color", None) if side is not None else None
            return getattr(c, "rgb", None)

        for name in self.profile["sheet_order"]:
            ws = self.wbf[name]
            first = self.srcf["header_row"] + 1 if name == C.SHEET_RAW else 3
            nc = self.raw_ncols if name == C.SHEET_RAW else ws.max_column
            for r in range(first, min(ws.max_row, first + 6)):
                for c in range(1, nc + 1):
                    b = ws.cell(r, c).border
                    coord = ws.cell(r, c).coordinate
                    if mode == "horizontal":
                        if st(b.left) or st(b.right):
                            bad.append((name, coord, "存在竖向边框"))
                        elif st(b.bottom) != "thin":
                            bad.append((name, coord, "缺水平浅灰边框"))
                        elif colr(b.bottom) != C.BORDER_COLOR:
                            bad.append((name, coord, "边框颜色非浅灰"))
                    else:
                        for side in (b.left, b.right, b.top, b.bottom):
                            if st(side) != "thin":
                                bad.append((name, coord, "四边框不完整"))
                                break
        self.add("R12", "边框模式=%s（%s）" % (mode, "仅浅灰水平线" if mode == "horizontal" else "四边浅灰"),
                 not bad, "不合规 %s" % bad[:6])

    # ---------- R13 冻结
    def r13_freeze(self):
        bad = [(n, self.wbf[n].freeze_panes) for n in self.profile["sheet_order"]
               if self.wbf[n].freeze_panes != C.FREEZE_PANES]
        self.add("R13", "全部工作表冻结表头首行（freeze=%s）" % C.FREEZE_PANES,
                 not bad, "不合规 %s" % bad)

    # ---------- R14 列宽
    def r14_width(self):
        bad = []
        for name in self.profile["sheet_order"]:
            ws = self.wbf[name]
            nc = self.raw_ncols if name == C.SHEET_RAW else ws.max_column
            for c in range(1, nc + 1):
                letter = get_column_letter(c)
                dim = ws.column_dimensions.get(letter)
                w = dim.width if dim is not None else None
                if w is None:
                    bad.append((name, letter, "未设置列宽"))
                elif not (C.COL_WIDTH_MIN - 0.01 <= w <= C.COL_WIDTH_MAX + 0.01):
                    bad.append((name, letter, w))
        self.add("R14", "全部列宽已自动适配（中文按全角计）且在 [%g,%g] 区间"
                 % (C.COL_WIDTH_MIN, C.COL_WIDTH_MAX), not bad, "不合规 %s" % bad[:8])

    # ---------- R15 标题行
    def r15_title(self):
        bad = []
        for name in self.profile["sheet_order"]:
            if name == C.SHEET_RAW and not self.profile["raw_override_title_row"]:
                continue  # spec：保留源第 1 行「单位：万元」，不覆盖
            ws = self.wbf[name]
            ncols = self.raw_ncols if name == C.SHEET_RAW else ws.max_column
            merged = any(str(m) == "A1:%s1" % get_column_letter(ncols) for m in ws.merged_cells.ranges)
            cell = ws["A1"]
            fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
            if not merged or fill != C.TITLE_FILL or cell.value != C.SHEET_TITLES[name]:
                bad.append((name, cell.value, merged, fill))
        self.add("R15", "各表第 1 行为跨列合并大标题（深蓝底白字加粗）", not bad,
                 "不合规 %s" % bad)

    # ---------- R16 条件格式
    def r16_cf(self):
        ws = self.wbf[C.SHEET_REASON]
        rules = [r for rng in ws.conditional_formatting for r in rng.rules]
        n = len(self.srcf["reasons"])
        if n == 0:
            self.add("R16", "未回款分类差异化柔和背景条件格式",
                     True, "源「未回款原因分类」字段 1000 行全为空，分类数=0，无数据行可上色（规则数=0）")
        else:
            self.add("R16", "未回款分类差异化柔和背景条件格式（%d 类 %d 条规则）" % (n, len(rules)),
                     len(rules) >= n, "规则数=%d 分类数=%d" % (len(rules), n))

    def run_all(self):
        for fn in (self.r1_sheets, self.r2_raw_intact, self.r3_formulas, self.r4_errors,
                   self.r5_rowcounts, self.r6_values, self.r7_order, self.r8_formats,
                   self.r9_align, self.r10_header, self.r11_zebra, self.r12_border,
                   self.r13_freeze, self.r14_width, self.r15_title, self.r16_cf):
            fn()
        return self.results


def main():
    p = argparse.ArgumentParser(description="结果文件交付前自检")
    p.add_argument("--file", required=True, help="Skill 产出的结果 xlsx")
    p.add_argument("--source", required=True, help="原始输入 xlsx")
    p.add_argument("--profile", default=C.DEFAULT_PROFILE, choices=sorted(C.PROFILES))
    p.add_argument("--json", default=None, help="把自检报告写成 JSON")
    args = p.parse_args()

    profile = C.get_profile(args.profile)
    ck = Checker(args.file, args.source, profile)
    results = ck.run_all()
    passed = sum(1 for r in results if r["ok"])
    report = {
        "file": os.path.abspath(args.file),
        "source": os.path.abspath(args.source),
        "profile": profile["name"],
        "total_checks": len(results),
        "passed": passed,
        "failed": len(results) - passed,
        "all_passed": passed == len(results),
        "items": results,
    }
    text = json.dumps(report, ensure_ascii=False, indent=2, default=str)
    if args.json:
        with open(args.json, "w", encoding="utf-8") as fh:
            fh.write(text)
    for r in results:
        print("%s %s %s" % ("PASS" if r["ok"] else "FAIL", r["id"], r["desc"]))
        if not r["ok"]:
            print("      %s" % r["detail"])
    print("\n合计 %d 项，通过 %d 项，未通过 %d 项 → %s"
          % (len(results), passed, len(results) - passed,
             "全部合规" if passed == len(results) else "存在不合规项"))
    sys.exit(0 if passed == len(results) else 1)


if __name__ == "__main__":
    main()
