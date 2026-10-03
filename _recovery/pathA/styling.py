# -*- coding: utf-8 -*-
"""
统一可视化美化模块（全工作表生效）。

规范来源：《题目描述》"统一可视化美化规范"，样式常量取自 config.py。
本模块只做样式，不写业务值，也不做任何统计计算。
"""

import unicodedata

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

import config as C


# ---------------------------------------------------------------- 工具

def display_width(text):
    """中文按全角字符标准计算显示宽度（全角=2，半角=1）。"""
    if text is None:
        return 0
    w = 0
    for ch in str(text):
        w += 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
    return w


def format_display(value, kind, money_format):
    """把值渲染成"单元格实际显示出来的样子"，仅用于估算列宽。"""
    if value is None or value == "":
        return ""
    if kind == "money":
        try:
            dec = money_format.count("0") - money_format.count("#")
            dec = max(dec, 0)
            return "{:,.{d}f}".format(float(value), d=dec)
        except (TypeError, ValueError):
            return str(value)
    if kind == "pct":
        try:
            return "{:.2f}%".format(float(value) * 100)
        except (TypeError, ValueError):
            return str(value)
    if kind == "int":
        try:
            return "{:,d}".format(int(round(float(value))))
        except (TypeError, ValueError):
            return str(value)
    return str(value)


def _fill(rgb):
    return PatternFill(fill_type="solid", start_color=rgb, end_color=rgb)


def _side(style, color):
    return Side(style=style, color=color) if style else Side()


# ---------------------------------------------------------------- 标题行 / 表头行

def write_title(ws, ncols, title):
    """第 1 行：跨全部列合并的大标题，深蓝底 + 白色加粗 + 居中。"""
    if not title:
        return
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    cell = ws.cell(1, 1)
    cell.value = title
    cell.font = Font(name=C.TITLE_FONT_NAME, size=C.TITLE_FONT_SIZE, bold=True,
                     color=C.TITLE_FONT_COLOR)
    cell.fill = _fill(C.TITLE_FILL)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # 合并区其余单元格也刷底色，避免出现色块断层
    for c in range(2, ncols + 1):
        cc = ws.cell(1, c)
        cc.fill = _fill(C.TITLE_FILL)
    ws.row_dimensions[1].height = C.TITLE_ROW_HEIGHT


def style_header(ws, ncols, header_row=2, set_alignment=True):
    """表头行：低饱和度蓝色背景 + 白色加粗字体 + 水平居中。"""
    fill = _fill(C.HEADER_FILL)
    for c in range(1, ncols + 1):
        cell = ws.cell(header_row, c)
        cell.font = Font(name=C.HEADER_FONT_NAME, size=C.HEADER_FONT_SIZE, bold=True,
                         color=C.HEADER_FONT_COLOR)
        cell.fill = fill
        if set_alignment:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = C.HEADER_ROW_HEIGHT


# ---------------------------------------------------------------- 数据区

def style_data(ws, ncols, first_row, nrows, kinds, align_by_kind, border_mode,
               number_formats=None, set_number_format=True, zebra=True,
               font_name=None, font_size=None):
    """
    数据区：斑马纹交替行底色 + 浅灰水平边框 + 按字段类型对齐。

    kinds: 每列的类型（int/money/pct/text/cat/date/num），长度 = ncols
    number_formats: 每列显式数字格式（优先于按 kind 推断）
    set_number_format=False 时完全不动数字格式（「原始数据」表用，保持源格式）
    """
    if nrows <= 0:
        return
    last_row = first_row + nrows - 1
    fname = font_name or C.DATA_FONT_NAME
    fsize = font_size or C.DATA_FONT_SIZE

    default_nf = {
        "money": (number_formats or {}).get("money", C.MONEY_FORMAT_SPEC),
        "int": C.INT_FORMAT,
        "pct": C.PCT_FORMAT,
        "date": C.DATE_FORMAT_SRC,
    }

    grey = C.BORDER_COLOR
    if border_mode == "all":
        border_data = Border(left=_side("thin", grey), right=_side("thin", grey),
                             top=_side("thin", grey), bottom=_side("thin", grey))
    else:  # horizontal：仅保留浅灰色水平边框
        border_data = Border(bottom=_side("thin", grey))

    zebra_fill = _fill(C.ZEBRA_FILL)
    no_fill = PatternFill(fill_type=None)

    for r in range(first_row, last_row + 1):
        striped = zebra and ((r - first_row) % 2 == 0)
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            kind = kinds[c - 1] if c - 1 < len(kinds) else "text"
            cell.font = Font(name=fname, size=fsize, color=C.DATA_FONT_COLOR)
            cell.alignment = Alignment(horizontal=align_by_kind.get(kind, "left"),
                                       vertical="center",
                                       wrap_text=(kind in ("text", "cat")))
            cell.border = border_data
            cell.fill = zebra_fill if striped else no_fill
            if set_number_format:
                if number_formats and c in number_formats:
                    cell.number_format = number_formats[c]
                elif kind in default_nf:
                    cell.number_format = default_nf[kind]


def header_border(ws, ncols, border_mode, header_row=2):
    """表头行边框：水平模式下补上下边框，形成清晰的表头分隔线。"""
    grey = C.BORDER_COLOR
    if border_mode == "all":
        b = Border(left=_side("thin", grey), right=_side("thin", grey),
                   top=_side("thin", grey), bottom=_side("thin", grey))
    else:
        b = Border(top=_side("thin", grey), bottom=_side("thin", grey))
    for c in range(1, ncols + 1):
        ws.cell(header_row, c).border = b


# ---------------------------------------------------------------- 列宽 / 冻结

def autofit(ws, ncols, headers, display_rows, kinds, money_format,
            skip_title_row=True):
    """
    全部列宽自动适配内容，中文按全角字符标准计算。

    display_rows 必须是"最终显示值"（不能是公式字符串），
    否则会把公式文本长度算进列宽——这正是附件2-示例列宽失真的原因。
    """
    for c in range(1, ncols + 1):
        kind = kinds[c - 1] if c - 1 < len(kinds) else "text"
        widths = [display_width(headers[c - 1]) if c - 1 < len(headers) else 0]
        for row in display_rows:
            if c - 1 < len(row):
                widths.append(display_width(format_display(row[c - 1], kind, money_format)))
        w = max(widths) + C.COL_WIDTH_PADDING if widths else C.COL_WIDTH_MIN
        ws.column_dimensions[get_column_letter(c)].width = max(
            C.COL_WIDTH_MIN, min(C.COL_WIDTH_MAX, float(w)))


def freeze(ws, panes=None):
    """冻结表头首行（标题行 + 表头行都在视图上方固定）。"""
    ws.freeze_panes = panes or C.FREEZE_PANES


# ---------------------------------------------------------------- 条件格式

def reason_conditional_formats(ws, col_idx, first_row, nrows, categories):
    """
    任务2：为不同未回款分类单元格设置差异化柔和背景条件格式。

    用真正的条件格式（FormulaRule）而非静态填充：分类文本命中即上色，
    数据行增删后规则依然生效，满足"动态"要求。
    """
    if nrows <= 0 or not categories:
        return 0
    letter = get_column_letter(col_idx)
    rng = "%s%d:%s%d" % (letter, first_row, letter, first_row + nrows - 1)
    n = 0
    for i, cat in enumerate(categories):
        color = C.REASON_SOFT_COLORS[i % len(C.REASON_SOFT_COLORS)]
        text = str(cat).replace('"', '""')
        rule = FormulaRule(formula=['EXACT($%s%d,"%s")' % (letter, first_row, text)],
                           fill=_fill(color), stopIfTrue=False)
        ws.conditional_formatting.add(rng, rule)
        n += 1
    return n
