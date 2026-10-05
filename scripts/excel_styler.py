# -*- coding: utf-8 -*-
"""
Excel 统一美化渲染模块 excel_styler

严格遵循题目可视化规范：
1、表头低饱和度蓝色背景，白色加粗居中；
2、斑马纹交替底色，仅保留浅灰色水平边框；
3、对齐：文本左对齐；数值右对齐；日期、分类居中；
4、自动列宽，中文等全角字符按东亚宽度（W/F）计 2、其余计 1；
5、冻结表头首行；开启 auto_filter 筛选；
6、数字格式：金额千分位 3 位小数；百分比 2 位小数；
7、支持任务2：未回款原因分类差异化柔和背景条件格式；
8、工作表第 1 行合并大标题。

性能要点
--------
openpyxl 的样式对象（Font / PatternFill / Alignment / Border）是**不可变、可共享**的：
本模块全部在循环外创建一次、循环内复用引用，避免逐格 new 样式对象
（千行级工作表上这是原先格式化阶段的主要开销）；
列宽按 `iter_rows` 单遍扫描估宽，不做逐列重复遍历。
"""
import unicodedata

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

from config_validator import validate_sheet_field_kinds

# 循环内复用的共享样式对象（不可变，可安全地赋给任意多个单元格）
_ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
_ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
_NO_FILL = PatternFill(fill_type=None)      # “无填充”，用于斑马纹的偶数行


def build_excel_style_from_config(style_cfg: dict) -> dict:
    """读取配置构建可复用的 openpyxl 样式对象集合（每次调用返回一套新对象）。"""
    header_fill = PatternFill(start_color=style_cfg["header_bg_color"],
                              end_color=style_cfg["header_bg_color"], fill_type="solid")
    header_font = Font(bold=True, color=style_cfg["header_font_color"])
    odd_row_fill = PatternFill(start_color=style_cfg["odd_row_bg"],
                               end_color=style_cfg["odd_row_bg"], fill_type="solid")
    thin_h_line = Side(style="thin", color=style_cfg["horizontal_border_color"])
    h_border = Border(top=None, bottom=thin_h_line, left=None, right=None)

    # 大标题样式
    title_fill = PatternFill(start_color=style_cfg["title_bg_color"],
                             end_color=style_cfg["title_bg_color"], fill_type="solid")
    title_font = Font(bold=True, color=style_cfg["title_font_color"],
                      size=style_cfg["title_font_size"])

    return {
        "header_fill": header_fill,
        "header_font": header_font,
        "odd_fill": odd_row_fill,
        "even_fill": _NO_FILL,          # 偶数行即“无填充”，共享同一对象
        "h_border": h_border,
        "title_fill": title_fill,
        "title_font": title_font,
    }


def set_sheet_title_row(ws: Worksheet, title_text: str, total_col_count: int) -> str:
    """
    设置工作表第 1 行合并大标题。
    :param ws: 工作表对象
    :param title_text: 标题显示文本
    :param total_col_count: 需要合并的总列数量
    :return: 合并区间字符串（形如 "A1:H1"）
    """
    end_letter = get_column_letter(total_col_count)
    merge_range = f"A1:{end_letter}1"
    ws.merge_cells(merge_range)
    ws["A1"].value = title_text
    # 标题的填充/字体/对齐由调用方（apply_sheet_format）统一设置
    return merge_range


def add_category_conditional_format(ws: Worksheet, cat_col_index: int, start_row: int,
                                    style_cfg: dict = None):
    """
    任务2专属：未回款原因分类列设置差异化柔和背景条件格式

    ⚠️ **默认不启用**（配置 `style_setting.reason_soft_cf` = false）：用户 2026-10-03 选定
    该列**不设条件格式**、与同行其他单元格一致（只走斑马纹）。启用时由
    `stat_engine` 传入 `enable_cat_cond_format=True` 呼叫本函数。

    ⚠️ 颜色必须是 **8 位 ARGB**（`FF......`）：openpyxl 会把 6 位十六进制补成
    `00RRGGBB`，alpha=00 即完全透明，条件格式看起来“没生效”（露出斑马纹底色）。

    ⚠️ 公式必须锚定**区间左上角行**：`EXACT($B3,"分类名")`。Excel 以 sqref 左上角为基准
    做**相对行偏移**——若第 k 条规则写成分类自己那一行（`$B4`、`$B5`…），判断第 4 行时
    它会被偏移成 `$B5`，于是**除首行外一条规则都不命中**，分类底色露出斑马纹。
    2026-10-03 修复：重构时曾把 `EXACT($B{first_row},…)` 写成分类自己所在行并丢失
    EXACT（原始正确实现见 `_recovery/pathA/styling.py`）。

    :param ws: 工作表对象
    :param cat_col_index: 分类列序号，从 1 开始
    :param start_row: 数据起始行（跳过表头）
    :param style_cfg: 样式配置；取 `reason_soft_colors`（缺省用内置 8 色）
    """
    style_cfg = style_cfg or {}
    soft_color_list = style_cfg.get("reason_soft_colors") or [
        "FFFDE9D9", "FFDCE6F1", "FFEBF1DE", "FFE4DFEC",
        "FFF2DCDB", "FFDAEEF3", "FFFFF2CC", "FFEAF1DD",
    ]
    max_row = ws.max_row
    col_letter = get_column_letter(cat_col_index)
    # 条件格式公式的锚点：区间左上角行（不是分类自己所在行）
    anchor_row = start_row
    range_str = f"{col_letter}{anchor_row}:{col_letter}{max_row}"
    used_values = set()
    for r in range(start_row, max_row + 1):
        val = ws.cell(row=r, column=cat_col_index).value
        if val is None or (isinstance(val, str) and not val.strip()) or val in used_values:
            continue
        used_values.add(val)
        # 第 1 个分类取第 1 个颜色（不用 len(...)%n，避免跳过首色）
        pick_color = soft_color_list[(len(used_values) - 1) % len(soft_color_list)]
        fill = PatternFill(start_color=pick_color, end_color=pick_color, fill_type="solid")
        rule = FormulaRule(
            formula=['EXACT($%s%d,"%s")' % (col_letter, anchor_row,
                                            str(val).replace('"', '""'))],
            fill=fill, stopIfTrue=True)
        ws.conditional_formatting.add(range_str, rule)


def display_width(text) -> int:
    """
    按**东亚宽度标准**计算显示宽度：`unicodedata.east_asian_width` 为 W（宽）/ F（全角）
    的字符计 2，其余（半角字母数字、ASCII 标点等）计 1。

    覆盖范围远大于"仅汉字区间 \u4e00-\u9fff"，还包括全角标点（，。：；「」（））、
    全角字母数字、假名、谚文等——否则「总合同金额（万元）」这类标题会被低估宽度而显示不全。
    """
    if text is None:
        return 0
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in str(text))


def format_display(value, number_format=None) -> str:
    """
    把单元格值渲染成"实际显示出来的样子"，仅用于估算列宽。

    例：903.192 配 `#,##0.000` → `903.192`；0.2543 配 `0.00%` → `25.43%`。
    直接用原始值会漏掉千分位与百分号占用的宽度。
    """
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        nf = (number_format or "General").strip()
        frac = nf.split(".", 1)[1] if "." in nf else ""
        dec = sum(1 for c in frac if c in "0#")
        if "%" in nf:
            return "{:.{d}f}%".format(float(value) * 100, d=dec)
        if nf in ("General", "@", ""):
            return str(value)
        if "," in nf:
            return "{:,.{d}f}".format(float(value), d=dec)
        return "{:.{d}f}".format(float(value), d=dec)
    return str(value)


def auto_fit_column_width(ws: Worksheet, style_cfg: dict = None, display_lookup: dict = None):
    """
    全部列宽按"最终显示出来的内容"自动适配：

    1. 中文等全角字符按 **2** 个字符宽度计（东亚宽度 W/F），半角字符按 1 计；
    2. 公式单元格取注入的**缓存值**（`display_lookup`，形如 {坐标: 值}）来估宽，
       而不是拿公式文本去量——否则长公式会把列宽撑爆；无缓存值时跳过该格；
    3. 列宽 = clamp(max(表头宽度, 各数据行显示宽度) + padding, min, max)，
       上下限与内边距由配置 `style_setting` 的
       `col_width_min` / `col_width_max` / `col_width_padding` 控制；
    4. 跳过合并单元格（大标题横跨全表，不能用来决定单列宽度）。

    实现用 `iter_rows` 对全表做**单遍**扫描、按列号累计最大宽度，
    替代原先「逐列取整列再逐格判断」的多遍写法。
    """
    style_cfg = style_cfg or {}
    lo = style_cfg.get("col_width_min", 8)
    hi = style_cfg.get("col_width_max", 40)
    pad = style_cfg.get("col_width_padding", 3)

    max_widths = {}                            # {列号: 该列最大显示宽度}
    for row in ws.iter_rows():
        for cell in row:
            if cell.coordinate in ws.merged_cells:
                continue
            if display_lookup and cell.coordinate in display_lookup:
                text = format_display(display_lookup[cell.coordinate], cell.number_format)
            else:
                val = cell.value
                if val is None:
                    continue
                if isinstance(val, str) and val.startswith("="):
                    continue          # 无缓存值可用时，不拿公式文本当宽度
                text = format_display(val, cell.number_format)
            w = display_width(text)
            if w > max_widths.get(cell.column, 0):
                max_widths[cell.column] = w

    for col_idx in range(1, ws.max_column + 1):
        w = max_widths.get(col_idx)
        final_w = (w + pad) if w is not None else lo
        ws.column_dimensions[get_column_letter(col_idx)].width = \
            max(lo, min(hi, float(final_w)))


def _looks_numeric(value) -> bool:
    """判断值是否“看起来是数值”：数字本身，或能被 float() 解析的字符串。"""
    if value is None or isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        t = value.strip()
        if not t:
            return False
        try:
            float(t)
            return True
        except ValueError:
            return False
    return False


def detect_numeric_text_columns(ws: Worksheet, header_row_idx: int, name2idx: dict,
                                text_names: list, min_samples: int = 3,
                                numeric_ratio: float = 0.8) -> list:
    """
    找出「登记为文本字段、但数据几乎全是数值」的列，供告警使用（不阻断生成）。

    这类列通常是**类型登记错了**：本该是数值列（右对齐 + 金额/百分比格式），
    却被放进了 text_fields，于是左对齐且不设数字格式。

    :return: [(表头, 列号, 数值占比), ...]
    """
    found = []
    if not text_names:
        return found
    for name in text_names:
        c_idx = name2idx.get(name)
        if not c_idx:
            continue
        vals = [ws.cell(r, c_idx).value for r in range(header_row_idx + 1, ws.max_row + 1)]
        vals = [v for v in vals if v not in (None, "")]
        if len(vals) < min_samples:
            continue
        hit = sum(1 for v in vals if _looks_numeric(v))
        ratio = hit / float(len(vals))
        if ratio >= numeric_ratio:
            found.append((name, c_idx, ratio))
    return found


def apply_sheet_format(ws: Worksheet, style_cfg: dict,
                       num_cols: list, pct_cols: list, date_cols: list, cat_cols: list,
                       header_row_idx: int = 2, enable_cat_cond_format: bool = False,
                       cat_field_name: str = "",
                       sheet_title: str = None, total_title_cols: int = 0,
                       int_cols: list = None, text_cols: list = None,
                       set_number_format: bool = True, display_values: dict = None):
    """
    对工作表执行全套美化。

    :param ws: worksheet 对象
    :param style_cfg: 样式配置字典
    :param num_cols: 金额字段名列表
    :param pct_cols: 百分比字段名列表
    :param date_cols: 日期字段列表
    :param cat_cols: 分类字段列表
    :param header_row_idx: 表头行号（新增大标题后，表头固定为第 2 行）
    :param enable_cat_cond_format: 是否开启分类差异化条件格式（任务2打开）
    :param cat_field_name: 条件格式绑定的分类字段名称
    :param sheet_title: 大标题文本，不为空则设置 A1 合并标题
    :param total_title_cols: 标题合并总列数（同时限定表头样式与筛选范围的有效列）
    :param int_cols: 整数字段名列表（右对齐 + 整数格式）
    :param text_cols: 文本字段名列表（显式左对齐）；所有列都必须在类型清单中登记
    :param set_number_format: False = 不改数字格式（「原始数据」等需逐格保持源格式的表）
    :param display_values: 公式格的缓存值 {坐标: 值}；列宽按"显示内容"估算时使用，
                          不传则公式格不参与列宽计算（只按表头与文本格估宽）
    """
    style_dict = build_excel_style_from_config(style_cfg)
    header_fill = style_dict["header_fill"]
    header_font = style_dict["header_font"]
    odd_fill = style_dict["odd_fill"]
    h_border = style_dict["h_border"]

    int_cols = int_cols or []
    text_cols = text_cols or []

    # ---- 第 1 行合并大标题 ----
    if sheet_title is not None and total_title_cols > 0:
        set_sheet_title_row(ws, sheet_title, total_title_cols)
        ws["A1"].fill = style_dict["title_fill"]
        ws["A1"].font = style_dict["title_font"]
        ws["A1"].alignment = _ALIGN_CENTER

    # 冻结窗格：大标题 1 行 + 表头 1 行，数据从第 3 行开始
    ws.freeze_panes = f"A{header_row_idx + 1}"

    # 表头名 → 列号（重复表头取最后一次出现，与原实现一致）
    name2idx = {}
    for idx, n in enumerate(ws[header_row_idx]):
        if n.value is not None:
            name2idx[n.value] = idx + 1

    # 字段类型清单：每一列都必须显式登记；有未登记的表头直接报错，
    # 避免新增列被静默地兜底成文本（左对齐、不设数字格式）。
    validate_sheet_field_kinds(
        ws.title, list(name2idx.keys()),
        {"num_fields": num_cols, "int_fields": int_cols, "pct_fields": pct_cols,
         "date_fields": date_cols, "cat_fields": cat_cols, "text_fields": text_cols},
    )

    # ---- 表头样式 + 筛选范围 ----
    # 用传入的 total_title_cols 限定有效列（源文件可能有带样式无表头的"幽灵列"）；
    # 未指定列数时兜底按整行处理。
    if total_title_cols > 0:
        for col_idx in range(1, total_title_cols + 1):
            cell = ws.cell(row=header_row_idx, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = _ALIGN_CENTER
        last_col_letter = get_column_letter(total_title_cols)
    else:
        for cell in ws[header_row_idx]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = _ALIGN_CENTER
        last_col_letter = get_column_letter(ws.max_column)
    # 筛选区域必须是标准 A1 记法，且严格限制在有效列内
    ws.auto_filter.ref = f"A{header_row_idx}:{last_col_letter}{ws.max_row}"

    # ---- 预计算每列的（列号, 对齐样式, 数字格式）动作表 ----
    # 对齐与数字格式只依赖「字段类型 × set_number_format 开关」，与行无关，
    # 全部在行循环外决定；共享样式对象直接复用（见模块头性能要点）。
    fmt_amount = style_cfg["amount_number_format"]
    fmt_int = style_cfg.get("int_number_format", "0")
    fmt_pct = style_cfg["percent_number_format"]
    fmt_date = style_cfg.get("date_number_format", "@")

    def _kind_action(kind: str):
        """字段类型 → (对齐样式, 数字格式)；数字格式 None 表示该列不改格式。"""
        if kind == "num":
            return _ALIGN_RIGHT, fmt_amount if set_number_format else None
        if kind == "int":
            return _ALIGN_RIGHT, fmt_int if set_number_format else None
        if kind == "pct":
            return _ALIGN_RIGHT, fmt_pct if set_number_format else None
        if kind == "date":
            # 日期字段居中；维持源文本录入格式（默认 @，不得转成日期序列值）
            return _ALIGN_CENTER, fmt_date if set_number_format else None
        if kind == "cat":
            return _ALIGN_CENTER, None
        return _ALIGN_LEFT, None      # text（未登记列已被上面的校验直接拦截）

    col_actions = []
    for col_name, c_idx in name2idx.items():
        if col_name in num_cols:
            kind = "num"
        elif col_name in int_cols:
            kind = "int"
        elif col_name in pct_cols:
            kind = "pct"
        elif col_name in date_cols:
            kind = "date"
        elif col_name in cat_cols:
            kind = "cat"
        else:
            kind = "text"             # 显式 text_fields 与兜底分支行为一致（左对齐）
        col_actions.append((c_idx,) + _kind_action(kind))

    # ---- 数据区：斑马纹（奇数数据行上底色，偶数行无填充）+ 浅灰水平边框 + 对齐/格式 ----
    for r in range(header_row_idx + 1, ws.max_row + 1):
        striped = ((r - header_row_idx) % 2 == 1)
        row_fill = odd_fill if striped else _NO_FILL
        for cell in ws[r]:
            cell.fill = row_fill
            cell.border = h_border
        for c_idx, align, nf in col_actions:
            cell = ws.cell(row=r, column=c_idx)
            cell.alignment = align
            if nf is not None:
                cell.number_format = nf

    # 任务2：未回款原因分类差异化条件格式
    if enable_cat_cond_format and cat_field_name in name2idx:
        add_category_conditional_format(ws, name2idx[cat_field_name],
                                        start_row=header_row_idx + 1, style_cfg=style_cfg)

    # 告警：被登记为文本字段、但数据几乎全是数值的列（很可能是类型登记错了）。
    # 只提示不阻断——真正的"未登记列"已在上面 validate_sheet_field_kinds 直接报错。
    for _name, _col, _ratio in detect_numeric_text_columns(ws, header_row_idx, name2idx,
                                                           text_cols):
        print("⚠️ [对齐告警]「%s」%s 列被登记为文本字段，但 %.0f%% 的数据是数值；"
              "如确为数值请移入 num_fields（右对齐 + 数字格式）"
              % (ws.title, get_column_letter(_col), _ratio * 100))

    auto_fit_column_width(ws, style_cfg, display_values)
