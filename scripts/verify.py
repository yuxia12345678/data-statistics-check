# -*- coding: utf-8 -*-
"""
交付前自检程序（独立复核，不复用生成侧 stat_engine 的任何计算结果）。

用法：
    python3 verify.py --file 结果.xlsx --source 附件1-合同开票及回款核对表.xlsx
    可选： --config config/contract_repayment.json
           --profile spec|example   --json report.json

设计原则：
  - 所有期望值都从"源文件"重新算一遍，再与结果文件的缓存值对照，
    形成源事实 → 产物 的独立验证链，不用生成逻辑自证。
  - 公式视图（data_only=False）核公式与样式；数值视图（data_only=True）核缓存值。
  - 任何一项不通过 → 退出码 1，并按题目"注意事项(4)"逐项列出。
"""

import argparse
import collections
import copy
import json
import re
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    sys.stdout.reconfigure(errors="replace")
except Exception:
    pass

# ---------------------------------------------------------------- 运行前依赖自检
# 放在 `from openpyxl import ...` 之前：齐全则跳过安装直接自检，缺失则离线优先补齐。
# 实现见 scripts/deps_check.py。
from deps_check import DependencyError, ensure_dependencies

try:
    ensure_dependencies()
except DependencyError as exc:
    print("❌ 运行前依赖自检未通过，已终止执行：\n      %s" % exc)
    sys.exit(1)

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


# ================================================================ 配置（唯一事实源）
# 产物生成与交付前自检共用同一份业务配置：config/contract_repayment.json。
# 本文件不再维护第二份常量，以下常量全部由该 JSON 派生，避免两份配置漂移。
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_CONFIG_PATH = os.path.normpath(
    os.path.join(_SCRIPT_DIR, "..", "config", "contract_repayment.json"))


class _Config:
    """`config/contract_repayment.json` 的只读常量视图（自检侧的唯一事实源）。"""

    def __init__(self, path):
        with open(path, "r", encoding="utf-8") as fh:
            self.raw = json.load(fh)
        cfg = self.raw
        style = cfg["style_setting"]
        names = cfg["output_sheet_names"]
        titles = cfg["sheet_titles"]
        t1 = cfg["task1_group_merge"]
        t2 = cfg["task2_special_agg"]
        t3 = cfg["task3_multi_dim"]

        # ---- 工作表名、顺序与表头大标题
        self.SHEET_RAW = names["raw_copy"]
        self.SHEET_MERGE = names["task1_result"]
        self.SHEET_REASON = names["task2"]
        self.SHEET_OVERVIEW = names["overview"]
        self.DIMENSIONS = [(d["sheet_name"], d["dimension_field"]) for d in t3["dim_list"]]
        self.SHEET_TITLES = {
            self.SHEET_RAW: titles["raw_copy"],
            self.SHEET_MERGE: titles["task1_result"],
            self.SHEET_REASON: titles["task2"],
            self.SHEET_OVERVIEW: titles["overview"],
        }
        for _dim_sheet, _dim_field in self.DIMENSIONS:
            self.SHEET_TITLES[_dim_sheet] = _dim_field + titles["dim_suffix"]

        # ---- 字段口径
        self.MERGE_SUM_FIELDS = [r["field_name"] for r in t1["agg_strategy_list"]
                                 if r["agg_strategy"] == "sum"]
        # 合并表「备注文本字段」的去重拼接分隔符（与生成侧同一事实源，题目只要求“统一”）
        self.JOIN_SEP = next((r.get("join_separator", "；") for r in t1["agg_strategy_list"]
                              if r["agg_strategy"] == "distinct_join"), "；")
        self.OVERVIEW_LABELS = [m["output_field"] for m in t3["overview_metrics"]]
        self.FIELD_KIND = dict(cfg["field_kind"])
        # 任务2「占比」分母定义：原始数据两列合计之差（被减数合计 − 减数合计）。
        # 与生成侧 stat_engine._resolve_ratio_denominator 共用同一份配置。
        self.T2_RATIO_DENOM = dict(t2.get("ratio_denominator") or {})

        # ---- 业务字段名（全部取自配置，代码中不出现任何业务字面量）----
        self.PRIMARY_KEY = t1["primary_key"]
        self.T2_GROUP_FIELD = t2["group_field"]
        self.T2_AMOUNT_FIELD = t2["amount_field"]
        self.T2_SORT_FIELD = t2.get("sort_by_field") or t2["amount_field"]
        self.T2_RATIO_FIELD = t2.get("ratio_field_name", "占比")

        # 维度表指标：求和项（保持 dim_metrics 声明顺序）+ 去重计数项 + 派生比率列表
        self.DIM_SUM_OUTPUTS = [k for k, v in t3["dim_metrics"].items() if v[1] == "sum"]
        self.DIM_SUM_SOURCE = {k: v[0] for k, v in t3["dim_metrics"].items()}
        self.DIM_SUM_SOURCE_ORDER = [self.DIM_SUM_SOURCE[k] for k in self.DIM_SUM_OUTPUTS]
        self.DIM_COUNT_OUTPUT = next((k for k, v in t3["dim_metrics"].items()
                                      if v[1] == "nunique"), None)
        self.DIM_COUNT_SOURCE = (t3["dim_metrics"][self.DIM_COUNT_OUTPUT][0]
                                 if self.DIM_COUNT_OUTPUT else self.PRIMARY_KEY)
        self.DIM_RATIOS = list(t3.get("derived_ratio_list") or [])
        # 维度表排序字段
        self.DIM_SORT_FIELD = t3.get("sort_by_field") or self.T2_SORT_FIELD

        # 统计总览指标（顺序即输出行顺序）
        self.OVERVIEW_METRICS = list(t3["overview_metrics"])

        # 「原始数据」中需要汇总合计的全部源字段（合并求和 + 维度求和 + 总览求和 + 占比分母）
        _totals = list(self.MERGE_SUM_FIELDS) + list(self.DIM_SUM_SOURCE_ORDER)
        for _m in self.OVERVIEW_METRICS:
            if _m.get("agg_operator") == "sum" and _m.get("source_field"):
                _totals.append(_m["source_field"])
        _totals.append(self.T2_AMOUNT_FIELD)
        _totals += [v for v in self.T2_RATIO_DENOM.values() if v]
        self.TOTAL_FIELDS = list(dict.fromkeys(_totals))

        # ---- 样式常量（全部取自 style_setting）
        self.TITLE_FILL = style["title_bg_color"]
        self.HEADER_FILL = style["header_bg_color"]
        self.HEADER_FONT_COLOR = style["header_font_color"]
        self.ZEBRA_FILL = style["odd_row_bg"]
        self.BORDER_COLOR = style["horizontal_border_color"]
        self.PCT_FORMAT = style["percent_number_format"]
        self.INT_FORMAT = style["int_number_format"]
        self.DATE_FORMAT_SRC = style["date_number_format"]
        self.COL_WIDTH_MIN = style["col_width_min"]
        self.COL_WIDTH_MAX = style["col_width_max"]
        self.FREEZE_PANES = style["freeze_panes"]
        # 「原始数据」是否叠加美化（false = 与源表逐格一致，样式类断言不适用于它）
        self.raw_beautify = bool(style.get("raw_beautify", False))
        # 任务2 分类列是否叠加「差异化柔和背景条件格式」
        # false = 该列不设条件格式（与同行其他单元格一致，只走斑马纹），R16/R17 断言规则数为 0
        self.REASON_SOFT_CF = bool(style.get("reason_soft_cf", False))

        # ---- profile（期望值口径）
        self.PROFILES = cfg["profiles"]
        self.DEFAULT_PROFILE = cfg["default_profile"]

    def get_profile(self, name=None):
        """返回 profile 的深拷贝，调用方可安全修改。"""
        key = name or self.DEFAULT_PROFILE
        if key not in self.PROFILES:
            raise ValueError("未知 profile：%s（可选 %s）"
                             % (key, "、".join(sorted(self.PROFILES))))
        return copy.deepcopy(self.PROFILES[key])

    def dim_metric_source(self, output_field):
        """维度指标输出名 → 其求和来源字段；「全局X」去掉「全局」前缀。"""
        key = output_field[2:] if output_field.startswith("全局") else output_field
        return self.DIM_SUM_SOURCE.get(key)

    def dim_sort_output(self):
        """维度表的排序输出列名（sort_by_field 指向的来源字段所对应的输出列）。"""
        for o in self.DIM_SUM_OUTPUTS:
            if self.DIM_SUM_SOURCE[o] == self.DIM_SORT_FIELD:
                return o
        return self.DIM_SUM_OUTPUTS[-1] if self.DIM_SUM_OUTPUTS else None


C = None    # 由 main() 依据 --config 构建；其余函数在运行期引用

TOL = 1e-6


# ================================================================ 源事实（独立重算）

def source_facts(path, profile=None, wb=None):
    """从源文件独立重算全部期望事实（不复用生成侧 stat_engine 的任何结果）。

    :param path: 源 xlsx 路径（wb 未传入时内部自行加载）
    :param profile: 期望值口径（example profile 的空值填充口径在此生效）
    :param wb: 可传入已加载的源工作簿复用，避免同一文件重复读盘
    """
    if wb is None:
        wb = load_workbook(path, data_only=True)
    ws = None
    for w in wb.worksheets:
        if w.sheet_state == "visible":
            ws = w
            break
    header_row = None
    for r in range(1, min(ws.max_row, 20) + 1):
        vals = [str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else ""
                for c in range(1, ws.max_column + 1)]
        if C.PRIMARY_KEY in vals:
            header_row = r
            break
    if header_row is None:
        raise ValueError("在源文件前 20 行内未找到业务主键字段「%s」，"
                         "请检查 --source 与配置 task1_group_merge.primary_key 是否一致"
                         % C.PRIMARY_KEY)
    col_of = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if v is not None and str(v).strip() and str(v).strip() not in col_of:
            col_of[str(v).strip()] = c

    raw = []          # [(row, {field: value})]
    for r in range(header_row + 1, ws.max_row + 1):
        rec = {}
        empty = True
        for f, c in col_of.items():
            v = ws.cell(r, c).value
            rec[f] = v
            if v not in (None, ""):
                empty = False
        if not empty:
            raw.append((r, rec))

    def num(v):
        if v in (None, ""):
            return 0.0
        try:
            return float(v)
        except (TypeError, ValueError):
            return 0.0

    facts = {
        "header_row": header_row,
        "col_of": col_of,
        "rows": len(raw),
        "raw_last_row": raw[-1][0] if raw else header_row,
        "totals": {f: sum(num(rec[f]) for _r, rec in raw)
                   for f in C.TOTAL_FIELDS if f in col_of},
        "raw": raw,
    }

    def key(rec, f):
        v = rec.get(f)
        return "" if v is None else (v if isinstance(v, str) else str(v)).strip()

    # 按业务主键分组
    g = collections.OrderedDict()
    for r, rec in raw:
        code = key(rec, C.PRIMARY_KEY)
        d = g.setdefault(code, {"sums": collections.defaultdict(float), "rows": 0,
                                "reasons": set()})
        d["rows"] += 1
        for f in C.MERGE_SUM_FIELDS:
            d["sums"][f] += num(rec.get(f))
        _r = key(rec, C.T2_GROUP_FIELD)
        if _r:
            d["reasons"].add(_r)
    facts["contracts"] = g
    facts["contract_count"] = len(g)

    # 维度分组
    dims = {}
    for _s, dim in C.DIMENSIONS:
        b = collections.OrderedDict()
        for r, rec in raw:
            k = key(rec, dim)
            if not k:
                continue
            d = b.setdefault(k, {"sums": collections.defaultdict(float), "codes": set(), "rows": 0})
            d["rows"] += 1
            d["codes"].add(key(rec, C.PRIMARY_KEY))
            for f in C.MERGE_SUM_FIELDS:
                d["sums"][f] += num(rec.get(f))
        dims[dim] = b
    facts["dims"] = dims

    # 分类字段（非空去重）
    # example profile 会把空值填充为「未填写」，源事实必须按同一口径重算，
    # 否则 R5/R6 会拿 spec 口径去核 example 产物。
    fill_blank = bool(profile and profile.get("fill_blank_reason"))
    b = collections.OrderedDict()
    for r, rec in raw:
        k = key(rec, C.T2_GROUP_FIELD)
        if not k:
            if not fill_blank:
                continue
            k = "未填写"
        d = b.setdefault(k, {"unpaid": 0.0, "codes": set(), "rows": 0})
        d["unpaid"] += num(rec.get(C.T2_AMOUNT_FIELD))
        d["rows"] += 1
        d["codes"].add(key(rec, C.PRIMARY_KEY))
    facts["reasons"] = b
    return facts


# ================================================================ 检查项

def close(a, b):
    if a is None or b is None:
        return a is None and b is None
    try:
        return abs(float(a) - float(b)) <= max(TOL, abs(float(b)) * 1e-9)
    except (TypeError, ValueError):
        return str(a) == str(b)


def _split_args(s):
    """按**顶层**逗号切分函数参数：忽略引号内、括号内的逗号。"""
    out, depth, cur, inq = [], 0, [], False
    for ch in s:
        if inq:
            cur.append(ch)
            if ch == '"':
                inq = False
        elif ch == '"':
            inq = True
            cur.append(ch)
        elif ch == "(":
            depth += 1
            cur.append(ch)
        elif ch == ")":
            depth -= 1
            cur.append(ch)
        elif ch == "," and depth == 0:
            out.append("".join(cur))
            cur = []
        else:
            cur.append(ch)
    out.append("".join(cur))
    return out


def iter_func_calls(formula, names):
    """逐个产出公式里的函数调用 `(函数名, [参数...])`，正确处理嵌套与引号。"""
    up = formula.upper()
    i = 0
    while True:
        best = None
        for nm in names:
            p = up.find(nm + "(", i)
            if p != -1 and (best is None or p < best[0]):
                best = (p, nm)
        if best is None:
            return
        pos, nm = best
        if pos > 0 and (up[pos - 1].isalnum() or up[pos - 1] in "_$!"):
            i = pos + len(nm) + 1
            continue
        open_idx = pos + len(nm)
        depth, k, inq = 0, open_idx, False
        while k < len(formula):
            ch = formula[k]
            if inq:
                if ch == '"':
                    inq = False
            elif ch == '"':
                inq = True
            elif ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
                if depth == 0:
                    break
            k += 1
        yield nm, _split_args(formula[open_idx + 1:k])
        i = k + 1


# 裸区间（无 `&""` 防护的整列/整块引用）
BARE_RANGE_RE = re.compile(r"^([^!]+!)?\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+$")
# 条件格式公式：EXACT($B3,"分类名")，锚点行 = 区间首行
CF_FORMULA_RE = re.compile(r'^EXACT\(\$([A-Z]{1,3})(\d+),"(.*)"\)$')


class Checker(object):
    def __init__(self, out_path, src_path, profile):
        # 完整工作表顺序（R1 用）；样式类断言（R9–R15）在 raw_beautify=false 时
        # 跳过「原始数据」——它与源表逐格一致，不参与统一美化。
        self.sheet_order_all = list(profile["sheet_order"])
        if not C.raw_beautify:
            profile["sheet_order"] = [s for s in profile["sheet_order"] if s != C.SHEET_RAW]
        self.profile = profile
        self.wbf = load_workbook(out_path, data_only=False)
        self.wbv = load_workbook(out_path, data_only=True)
        # 源工作簿只读盘一次：既供 source_facts 独立重算，也供 R2 的逐格比对使用
        self.src_wb = load_workbook(src_path, data_only=True)
        self.srcf = source_facts(src_path, profile, wb=self.src_wb)
        # 源文件存在带样式但无表头/无数据的"幽灵列"（本例 P..S），
        # 属于源文件自带状态；「原始数据」必须原样保留，故美化与核验都只覆盖有效列。
        self.raw_ncols = max(self.srcf["col_of"].values())
        self.results = []

    def close(self):
        """释放全部工作簿句柄（Windows 下避免影响后续流程 / 网盘同步）。"""
        for wb in (self.wbf, self.wbv, self.src_wb):
            try:
                wb.close()
            except Exception:
                pass

    def add(self, rid, desc, ok, detail=""):
        self.results.append({"id": rid, "desc": desc, "ok": bool(ok), "detail": detail})

    # ---------- R1 工作表
    def r1_sheets(self):
        want = self.sheet_order_all
        got = self.wbf.sheetnames
        self.add("R1", "8 张工作表名称与顺序严格匹配", got == want,
                 "期望=%s 实际=%s" % (want, got))

    # ---------- R2 原始数据原样
    def r2_raw_intact(self):
        """
        「原始数据」= 源表 Sheet1 的原样副本（仅改名），因此**全表逐格**比对：
        值、数字格式、填充/字体/边框/对齐、行数/列数、列宽、冻结窗格。
        """
        ws = self.wbv[C.SHEET_RAW]
        wsf = self.wbf[C.SHEET_RAW]
        src_ws = None
        for w in self.src_wb.worksheets:
            if w.sheet_state == "visible":
                src_ws = w
                break

        def _color(c):
            try:
                v = getattr(c, "rgb", None)
                return v if isinstance(v, str) else None
            except Exception:
                return None

        def cell_style(cell):
            f, fill, b, a = cell.font, cell.fill, cell.border, cell.alignment
            return ((fill.fill_type if fill else None,
                     _color(fill.start_color) if fill else None),
                    (f.bold, f.italic, f.size, f.name, _color(f.color)),
                    (getattr(b.left, "style", None), getattr(b.right, "style", None),
                     getattr(b.top, "style", None), getattr(b.bottom, "style", None)),
                    (a.horizontal, a.vertical, a.wrap_text))

        def width_of(sheet, idx):
            return sheet.column_dimensions[get_column_letter(idx)].width

        max_r = max(src_ws.max_row, ws.max_row)
        max_c = max(src_ws.max_column, ws.max_column)
        diff_val, diff_nf, diff_style, n = [], [], [], 0
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                n += 1
                sc, gc = src_ws.cell(r, c), ws.cell(r, c)
                v, gv = sc.value, gc.value
                if isinstance(v, (int, float)) and isinstance(gv, (int, float)):
                    if not close(v, gv):
                        diff_val.append((r, c, v, gv))
                elif (v if v is not None else "") != (gv if gv is not None else ""):
                    diff_val.append((r, c, repr(v)[:40], repr(gv)[:40]))
                if sc.number_format != gc.number_format:
                    diff_nf.append((r, c, sc.number_format, gc.number_format))
                if cell_style(sc) != cell_style(wsf.cell(r, c)):
                    diff_style.append((r, c))
        self.add("R2a", "「原始数据」全部单元格值与源文件一致（含行列顺序）",
                 not diff_val, "比对 %d 格，差异 %d 处 %s" % (n, len(diff_val), diff_val[:5]))
        self.add("R2b", "「原始数据」数字格式 + 填充/字体/边框/对齐 逐格未被修改",
                 not diff_nf and not diff_style,
                 "数字格式差异 %d 处 %s；样式差异 %d 处 %s"
                 % (len(diff_nf), diff_nf[:3], len(diff_style), diff_style[:5]))
        hr = self.srcf["header_row"]
        heads_ok = all(ws.cell(hr, c).value == src_ws.cell(hr, c).value
                       for c in range(1, max_c + 1))
        self.add("R2c", "「原始数据」表头文字与行位置未变", heads_ok,
                 "表头行=%d" % hr)
        wdiff = []
        for c in range(1, max_c + 1):
            if width_of(src_ws, c) != width_of(ws, c):
                wdiff.append((get_column_letter(c), width_of(src_ws, c), width_of(ws, c)))
        self.add("R2d", "「原始数据」行数/列数/列宽/冻结窗格与源一致",
                 (src_ws.max_row == ws.max_row and src_ws.max_column == ws.max_column
                  and not wdiff and src_ws.freeze_panes == ws.freeze_panes),
                 "行 源=%d 结果=%d；列 源=%d 结果=%d；列宽差异 %d %s；冻结 源=%r 结果=%r"
                 % (src_ws.max_row, ws.max_row, src_ws.max_column, ws.max_column,
                    len(wdiff), wdiff[:5], src_ws.freeze_panes, ws.freeze_panes))

    # ---------- R3 禁止硬编码
    def r3_formulas(self):
        """R3a 数值格必须是公式；R3b 公式格必须带缓存值；R3c 公式必须引用数据源。

        前两项在**同一次**全表扫描里完成：公式视图（wbf）看公式文本，
        数值视图（wbv）看缓存值，避免对同一批单元格扫两遍。
        """
        stat_sheets = [s for s in self.profile["sheet_order"] if s != C.SHEET_RAW]
        no_cache, numeric_hardcoded = [], []
        for name in stat_sheets:
            wsf, wsv = self.wbf[name], self.wbv[name]
            for r in range(3, wsf.max_row + 1):
                for c in range(1, wsf.max_column + 1):
                    f = wsf.cell(r, c).value
                    cached = wsv.cell(r, c).value
                    is_formula = isinstance(f, str) and f.startswith("=")
                    # R3b：公式格在缓存视图里必须有值；纯文本键值（维度名/合同号/
                    # 指标名等标识符，允许为字面量）与空格不算缺缓存
                    if cached is None and f not in (None, "") \
                            and (not isinstance(f, str) or is_formula):
                        no_cache.append((name, wsf.cell(r, c).coordinate))
                    # R3a：缓存值是数值的格子必须是公式（禁止硬编码数值）
                    if isinstance(cached, (int, float)) and not is_formula:
                        numeric_hardcoded.append((name, wsf.cell(r, c).coordinate, cached))
        self.add("R3a", "全部数值单元格均为动态公式（禁止硬编码）",
                 not numeric_hardcoded, "硬编码 %d 处 %s" % (len(numeric_hardcoded),
                                                          numeric_hardcoded[:6]))
        self.add("R3b", "公式单元格均带缓存值（程序化读取不为空）", not no_cache,
                 "缺缓存 %d 处 %s" % (len(no_cache), no_cache[:6]))
        ok = self._all_formulas_reference_raw(stat_sheets)
        self.add("R3c", "公式全部动态引用「原始数据」区间（或引用同表公式格派生）",
                 ok, "存在未引用数据源的公式：%s" % (getattr(self, "_bad_formula", None),))

    # 不含数据源引用的公式，只允许两种形态：
    #   1) 行号序号          ROW()-2
    #   2) 同表派生比率       IF(C3=0,"",E3/C3)  ← 分子分母本身都是引用数据源的公式格
    LOCAL_OK = (re.compile(r"^ROW\(\)-\d+$"),
                re.compile(r'^IF\(\$?[A-Z]{1,3}\$?\d+=0,"",\$?[A-Z]{1,3}\$?\d+/\$?[A-Z]{1,3}\$?\d+\)$'))

    def _all_formulas_reference_raw(self, stat_sheets):
        for name in stat_sheets:
            wsf = self.wbf[name]
            for row in wsf.iter_rows(min_row=3):
                for cell in row:
                    v = cell.value
                    if not (isinstance(v, str) and v.startswith("=")):
                        continue
                    body = v[1:]
                    if C.SHEET_RAW in body:
                        continue
                    if not any(p.match(body) for p in self.LOCAL_OK):
                        self._bad_formula = (name, cell.coordinate, body[:60])
                        return False
        return True

    # ---------- R4 公式错误
    def r4_errors(self):
        err_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!")
        # 空值残留：源表空单元格经 pandas 读入为 NaN，若判空失真会写出字面量 "nan"
        miss_tokens = ("nan", "nat", "none", "inf", "-inf")
        found = []
        leaked = []
        for name in self.wbv.sheetnames:
            ws = self.wbv[name]
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if not isinstance(v, str):
                        continue
                    if v in err_tokens:
                        found.append((name, cell.coordinate, v))
                    elif any(tok.strip().lower() in miss_tokens for tok in v.split("、")):
                        leaked.append((name, cell.coordinate, v))
        ok = not found and not leaked
        self.add("R4", "无公式错误值 / 无空值残留（nan·NaT 未被当成内容写入）", ok,
                 "错误 %d 处 %s；空值残留 %d 处 %s"
                 % (len(found), found[:4], len(leaked), leaked[:4]))

    # ---------- R5 行数
    def r5_rowcounts(self):
        f = self.srcf
        ws = self.wbv[C.SHEET_MERGE]
        self.add("R5a", "合并表行数 = 不同%s数" % C.PRIMARY_KEY,
                 ws.max_row - 2 == f["contract_count"],
                 "结果=%d 期望=%d" % (ws.max_row - 2, f["contract_count"]))
        for sheet, dim in C.DIMENSIONS:
            ws = self.wbv[sheet]
            self.add("R5-" + dim, "%s 行数 = %s 唯一值数" % (sheet, dim),
                     ws.max_row - 2 == len(f["dims"][dim]),
                     "结果=%d 期望=%d" % (ws.max_row - 2, len(f["dims"][dim])))
        ws = self.wbv[C.SHEET_OVERVIEW]
        n_ov = len(C.OVERVIEW_METRICS)
        self.add("R5-总览", "统计总览为 %d 项全局指标" % n_ov, ws.max_row - 2 == n_ov,
                 "结果=%d 期望=%d" % (ws.max_row - 2, n_ov))
        ws = self.wbv[C.SHEET_REASON]
        self.add("R5-原因", "%s 汇总表行数 = 非空去重分类数" % C.T2_GROUP_FIELD,
                 ws.max_row - 2 == len(f["reasons"]),
                 "结果=%d 期望=%d（源字段全为空时为 0）" % (ws.max_row - 2, len(f["reasons"])))

    # ---------- R6 数值正确
    def r6_values(self):
        f = self.srcf
        pk = C.PRIMARY_KEY
        # 合并表：按业务主键核对全部求和字段
        ws = self.wbv[C.SHEET_MERGE]
        headers = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        idx = {h: i + 1 for i, h in enumerate(headers)}
        bad = []
        bad_txt = []
        for r in range(3, ws.max_row + 1):
            code = ws.cell(r, idx[pk]).value
            g = f["contracts"].get(str(code).strip())
            if g is None:
                bad.append((r, code, "源中不存在该%s" % pk))
                continue
            for fld in C.MERGE_SUM_FIELDS:
                if fld not in idx:
                    continue
                if not close(ws.cell(r, idx[fld]).value, g["sums"][fld]):
                    bad.append((r, code, fld, ws.cell(r, idx[fld]).value, g["sums"][fld]))
            # 备注文本字段：按配置分隔符去重拼接（题目要求“统一分隔符”，集合口径比对）
            if C.T2_GROUP_FIELD in idx:
                _v = ws.cell(r, idx[C.T2_GROUP_FIELD]).value
                txt = "" if _v is None else str(_v).strip()
                got = {t for t in txt.split(C.JOIN_SEP) if t} if txt else set()
                if got != g["reasons"]:
                    bad_txt.append((r, code, txt, sorted(g["reasons"])))
        self.add("R6a", "合并表 %d 个数值字段求和正确 + 「%s」按「%s」去重拼接与源一致（%d 个%s）"
                 % (len(C.MERGE_SUM_FIELDS), C.T2_GROUP_FIELD, C.JOIN_SEP,
                    f["contract_count"], pk),
                 not bad and not bad_txt,
                 "数值差异 %d 处 %s；拼接差异 %d 处 %s"
                 % (len(bad), bad[:5], len(bad_txt), bad_txt[:3]))

        # 维度表：按配置的 dim_metrics / derived_ratio_list 逐列核对
        for sheet, dim in C.DIMENSIONS:
            ws = self.wbv[sheet]
            headers = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
            n_metrics = max(ws.max_column - 2, 0)
            bad = []
            for r in range(3, ws.max_row + 1):
                k = str(ws.cell(r, 2).value).strip()
                d = f["dims"][dim].get(k)
                if d is None:
                    bad.append((r, k, "源中无此维度值"))
                    continue
                exp = [d["sums"][src] for src in C.DIM_SUM_SOURCE_ORDER]
                exp.append(len(d["codes"]))
                for ratio in C.DIM_RATIOS:
                    num_src = C.dim_metric_source(ratio.get("numerator_field", ""))
                    _den_out = ratio.get("denominator_field", "")
                    if _den_out.startswith("全局"):
                        den_src = C.dim_metric_source(_den_out)
                        den = f["totals"].get(den_src)
                    else:
                        den_src = C.dim_metric_source(_den_out)
                        den = d["sums"].get(den_src)
                    nv = d["sums"].get(num_src) if num_src else None
                    exp.append((nv / den) if (nv is not None and den) else None)
                got = [ws.cell(r, c).value for c in range(3, 3 + len(exp))]
                for i, (a, b) in enumerate(zip(got, exp)):
                    if not close(a, b):
                        bad.append((sheet, r, k, headers[i + 2] if i + 2 < len(headers) else i + 2, a, b))
                if n_metrics != len(exp):
                    bad.append((sheet, "列数不匹配", n_metrics, len(exp)))
            self.add("R6-" + dim, "%s 全部 %d 项指标与源事实一致" % (sheet, len(exp)), not bad,
                     "差异 %d 处 %s" % (len(bad), bad[:5]))

        # 统计总览：按 overview_metrics 逐项重算
        ws = self.wbv[C.SHEET_OVERVIEW]
        exp = []
        for m in C.OVERVIEW_METRICS:
            op = m.get("agg_operator")
            src = m.get("source_field") or ""
            if op == "sum":
                exp.append(f["totals"].get(src))
            elif op == "nunique":
                if src == C.PRIMARY_KEY:
                    exp.append(f["contract_count"])
                elif src == C.T2_GROUP_FIELD:
                    exp.append(len(f["reasons"]))
                else:
                    exp.append(len({str(rec.get(src)).strip()
                                    for _r, rec in f["raw"] if rec.get(src) not in (None, "")}))
            elif op == "count_rows":
                exp.append(f["rows"])
            elif op == "ratio":
                _num = next((x["output_field"] for x in C.OVERVIEW_METRICS
                             if x.get("output_field") == m.get("numerator")), None)
                _den = next((x["output_field"] for x in C.OVERVIEW_METRICS
                             if x.get("output_field") == m.get("denominator")), None)
                _i_num = C.OVERVIEW_LABELS.index(_num) if _num in C.OVERVIEW_LABELS else None
                _i_den = C.OVERVIEW_LABELS.index(_den) if _den in C.OVERVIEW_LABELS else None
                nv = exp[_i_num] if _i_num is not None and _i_num < len(exp) else None
                dv = exp[_i_den] if _i_den is not None and _i_den < len(exp) else None
                exp.append((nv / dv) if (nv is not None and dv) else None)
            else:
                exp.append(None)
        got = [ws.cell(3 + i, 2).value for i in range(len(C.OVERVIEW_LABELS))]
        bad = [(C.OVERVIEW_LABELS[i], got[i], exp[i]) for i in range(len(exp))
               if not close(got[i], exp[i])]
        self.add("R6-总览", "统计总览 %d 项指标与源事实一致" % len(exp), not bad, "差异 %s" % bad)

        # 分类表
        ws = self.wbv[C.SHEET_REASON]
        bad = []
        # 占比分母 = 原始数据「被减数」合计 −「减数」合计
        # （口径来自配置 task2_special_agg.ratio_denominator，两字段直接取自数据源）
        _dn = C.T2_RATIO_DENOM
        _minuend = f["totals"].get(_dn.get("minuend_field", ""), 0.0)
        _subtrahend = f["totals"].get(_dn.get("subtrahend_field", ""), 0.0)
        reason_denom = _minuend - _subtrahend
        for r in range(3, ws.max_row + 1):
            k = str(ws.cell(r, 2).value).strip()
            d = f["reasons"].get(k)
            if d is None:
                bad.append((r, k, "源中无此分类"))
                continue
            for col, e in ((3, d["unpaid"]), (4, len(d["codes"])), (5, d["rows"]),
                           (6, d["unpaid"] / reason_denom if reason_denom else None)):
                if not close(ws.cell(r, col).value, e):
                    bad.append((r, k, col, ws.cell(r, col).value, e))
        self.add("R6-原因", "%s 汇总 4 项指标与源事实一致（占比分母 =「%s」合计 −「%s」合计）"
                 % (C.T2_GROUP_FIELD, _dn.get("minuend_field", ""), _dn.get("subtrahend_field", "")),
                 not bad, "差异 %s（源字段全空时本项为空表，自动通过）" % bad[:5])

    # ---------- R7 排序
    def r7_order(self):
        for sheet, dim in C.DIMENSIONS:
            ws = self.wbv[sheet]
            heads = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
            _sort_out = C.dim_sort_output()
            sc = (heads.index(_sort_out) + 1) if _sort_out in heads else 6
            vals = [ws.cell(r, sc).value or 0 for r in range(3, ws.max_row + 1)]
            self.add("R7-" + dim, "%s 按%s降序" % (sheet, C.DIM_SORT_FIELD),
                     all(vals[i] >= vals[i + 1] for i in range(len(vals) - 1)),
                     "首尾=%s/%s" % (vals[:1], vals[-1:]))
        ws = self.wbv[C.SHEET_REASON]
        heads = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        sc = (heads.index(C.T2_SORT_FIELD) + 1) if C.T2_SORT_FIELD in heads else 3
        vals = [ws.cell(r, sc).value or 0 for r in range(3, ws.max_row + 1)]
        self.add("R7-原因", "%s 汇总按%s降序" % (C.SHEET_REASON, C.T2_SORT_FIELD),
                 all(vals[i] >= vals[i + 1] for i in range(len(vals) - 1)), "行数=%d" % len(vals))
        ws = self.wbv[C.SHEET_MERGE]
        heads = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        key_col = heads.index(C.PRIMARY_KEY) + 1
        vals = [str(ws.cell(r, key_col).value or "") for r in range(3, ws.max_row + 1)]
        self.add("R7-合并", "合并表按%s升序（从小到大）" % C.PRIMARY_KEY,
                 all(vals[i] <= vals[i + 1] for i in range(len(vals) - 1)),
                 "行数=%d" % len(vals))
        # 序号连续
        for name in [C.SHEET_MERGE, C.SHEET_REASON] + [s for s, _d in C.DIMENSIONS]:
            ws = self.wbv[name]
            seq = [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)]
            self.add("R7seq-" + name, "%s 序号全新且连续（1..N）" % name,
                     seq == list(range(1, len(seq) + 1)), "前3=%s 末=%s" % (seq[:3], seq[-1:]))

    # ---------- R8 数字格式
    def r8_formats(self):
        money = self.profile["money_format"]
        bad = []
        for name in self.profile["sheet_order"]:
            if name == C.SHEET_RAW:
                continue
            ws = self.wbf[name]
            headers = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
            for c, h in enumerate(headers, start=1):
                kind = C.FIELD_KIND.get(str(h).replace("总", "") if h else "", None) or \
                       C.FIELD_KIND.get(h, "text")
                if name == C.SHEET_OVERVIEW and c == 2:
                    continue  # 逐行类型不同，单独核
                want = {"money": money, "pct": C.PCT_FORMAT, "int": C.INT_FORMAT,
                        "date": C.DATE_FORMAT_SRC}.get(kind)
                if not want:
                    continue
                for r in range(3, ws.max_row + 1):
                    nf = ws.cell(r, c).number_format
                    if nf != want:
                        bad.append((name, ws.cell(r, c).coordinate, nf, want))
                        break
        ws = self.wbf[C.SHEET_OVERVIEW]
        want_by_row = [money, money, money, money, C.PCT_FORMAT, C.INT_FORMAT,
                       C.INT_FORMAT, C.INT_FORMAT]
        for i, w in enumerate(want_by_row):
            nf = ws.cell(3 + i, 2).number_format
            if nf != w:
                bad.append((C.SHEET_OVERVIEW, ws.cell(3 + i, 2).coordinate, nf, w))
        dec = len(money.split(".")[1]) if "." in money else 0
        self.add("R8", "金额千分位+%d位小数、百分比2位小数、整数与日期格式合规"
                 % dec, not bad, "不合规 %s" % bad[:8])

    # ---------- R9 对齐
    def r9_align(self):
        align = self.profile["align_by_kind"]
        bad = []
        for name in self.profile["sheet_order"]:
            ws = self.wbf[name]
            hr = self.srcf["header_row"] if name == C.SHEET_RAW else 2
            nc = self.raw_ncols if name == C.SHEET_RAW else ws.max_column
            headers = [ws.cell(hr, c).value for c in range(1, nc + 1)]
            for c, h in enumerate(headers, start=1):
                kind = C.FIELD_KIND.get(h, "text")
                if name == C.SHEET_OVERVIEW:
                    kind = "text" if c == 1 else "num"
                want = align.get(kind, "left")
                for r in range(hr + 1, ws.max_row + 1):
                    got = ws.cell(r, c).alignment.horizontal
                    if got != want:
                        bad.append((name, ws.cell(r, c).coordinate, kind, got, want))
                        break
        self.add("R9", "文本左对齐 / 数值右对齐 / 日期与分类居中", not bad,
                 "不合规 %s" % bad[:8])

    # ---------- R10 表头样式
    def r10_header(self):
        bad = []
        for name in self.profile["sheet_order"]:
            ws = self.wbf[name]
            hr = self.srcf["header_row"] if name == C.SHEET_RAW else 2
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(hr, c)
                if cell.value in (None, ""):
                    continue
                fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
                colr = cell.font.color.rgb if cell.font and cell.font.color else None
                if fill != C.HEADER_FILL or not cell.font.b or colr != C.HEADER_FONT_COLOR \
                        or cell.alignment.horizontal != "center":
                    bad.append((name, cell.coordinate, fill, cell.font.b, colr,
                                cell.alignment.horizontal))
        self.add("R10", "表头低饱和度蓝底 + 白色加粗 + 水平居中", not bad,
                 "不合规 %s" % bad[:6])

    # ---------- R11 斑马纹
    def r11_zebra(self):
        """
        数据区斑马纹：奇数数据行 = 斑马纹底色，偶数数据行 = 无底色。
        2026-10-03 升级为**逐格**比对（原先只抽查第 1 列）——这样任何单列被单独涂色
        都会露馅（任务2 分类列曾叠加深色条件格式，在 Excel 里看成“与同行不一致”）。
        """
        bad = []
        for name in self.profile["sheet_order"]:
            ws = self.wbf[name]
            first = self.srcf["header_row"] + 1 if name == C.SHEET_RAW else 3
            nc = self.raw_ncols if name == C.SHEET_RAW else ws.max_column
            for i, r in enumerate(range(first, ws.max_row + 1)):
                want = C.ZEBRA_FILL if i % 2 == 0 else None
                for c in range(1, nc + 1):
                    fill = ws.cell(r, c).fill
                    rgb = fill.fgColor.rgb if fill and fill.fill_type else None
                    if rgb != want:
                        bad.append((name, ws.cell(r, c).coordinate, rgb, want))
        self.add("R11", "数据区斑马纹交替行底色（逐格，含分类列）", not bad,
                 "不合规 %s" % bad[:6])

    # ---------- R12 边框
    def r12_border(self):
        mode = self.profile["border_mode"]
        bad = []

        def st(side):
            return getattr(side, "style", None) if side is not None else None

        def colr(side):
            c = getattr(side, "color", None) if side is not None else None
            return getattr(c, "rgb", None)

        for name in self.profile["sheet_order"]:
            ws = self.wbf[name]
            first = self.srcf["header_row"] + 1 if name == C.SHEET_RAW else 3
            nc = self.raw_ncols if name == C.SHEET_RAW else ws.max_column
            for r in range(first, min(ws.max_row, first + 6)):
                for c in range(1, nc + 1):
                    b = ws.cell(r, c).border
                    coord = ws.cell(r, c).coordinate
                    if mode == "horizontal":
                        if st(b.left) or st(b.right):
                            bad.append((name, coord, "存在竖向边框"))
                        elif st(b.bottom) != "thin":
                            bad.append((name, coord, "缺水平浅灰边框"))
                        elif colr(b.bottom) != C.BORDER_COLOR:
                            bad.append((name, coord, "边框颜色非浅灰"))
                    else:
                        for side in (b.left, b.right, b.top, b.bottom):
                            if st(side) != "thin":
                                bad.append((name, coord, "四边框不完整"))
                                break
        self.add("R12", "边框模式=%s（%s）" % (mode, "仅浅灰水平线" if mode == "horizontal" else "四边浅灰"),
                 not bad, "不合规 %s" % bad[:6])

    # ---------- R13 冻结
    def r13_freeze(self):
        bad = [(n, self.wbf[n].freeze_panes) for n in self.profile["sheet_order"]
               if self.wbf[n].freeze_panes != C.FREEZE_PANES]
        self.add("R13", "全部工作表冻结表头首行（freeze=%s）" % C.FREEZE_PANES,
                 not bad, "不合规 %s" % bad)

    # ---------- R14 列宽
    def r14_width(self):
        bad = []
        for name in self.profile["sheet_order"]:
            ws = self.wbf[name]
            nc = self.raw_ncols if name == C.SHEET_RAW else ws.max_column
            for c in range(1, nc + 1):
                letter = get_column_letter(c)
                dim = ws.column_dimensions.get(letter)
                w = dim.width if dim is not None else None
                if w is None:
                    bad.append((name, letter, "未设置列宽"))
                elif not (C.COL_WIDTH_MIN - 0.01 <= w <= C.COL_WIDTH_MAX + 0.01):
                    bad.append((name, letter, w))
        self.add("R14", "全部列宽已自动适配（中文按全角计）且在 [%g,%g] 区间"
                 % (C.COL_WIDTH_MIN, C.COL_WIDTH_MAX), not bad, "不合规 %s" % bad[:8])

    # ---------- R15 标题行
    def r15_title(self):
        bad = []
        for name in self.profile["sheet_order"]:
            if name == C.SHEET_RAW and not self.profile["raw_override_title_row"]:
                continue  # spec：保留源第 1 行「单位：万元」，不覆盖
            ws = self.wbf[name]
            ncols = self.raw_ncols if name == C.SHEET_RAW else ws.max_column
            merged = any(str(m) == "A1:%s1" % get_column_letter(ncols) for m in ws.merged_cells.ranges)
            cell = ws["A1"]
            fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
            if not merged or fill != C.TITLE_FILL or cell.value != C.SHEET_TITLES[name]:
                bad.append((name, cell.value, merged, fill))
        self.add("R15", "各表第 1 行为跨列合并大标题（深蓝底白字加粗）", not bad,
                 "不合规 %s" % bad)

    # ---------- R16 条件格式
    def r16_cf(self):
        ws = self.wbf[C.SHEET_REASON]
        rules = [r for rng in ws.conditional_formatting for r in rng.rules]
        n = len(self.srcf["reasons"])
        if not C.REASON_SOFT_CF:
            # 配置关闭：该列不设条件格式，与同行其他单元格一样只走斑马纹（见 R11）。
            self.add("R16", "%s 柔和背景条件格式：配置已关闭（reason_soft_cf=false）" % C.T2_GROUP_FIELD,
                     len(rules) == 0,
                     "规则数应为 0（实际 %d）；该列底色应等于同行斑马纹" % len(rules))
            return
        # 柔和背景必须是**8 位不透明 ARGB**：6 位十六进制会被 openpyxl 补成 00RRGGBB，
        # alpha=00 即透明 —— 条件格式形同失效（露出斑马纹底色），属缺陷。
        fills = []
        for rule in rules:
            f = rule.dxf.fill if rule.dxf else None
            try:
                rgb = str(f.start_color.rgb) if f is not None and f.start_color is not None else None
            except Exception:
                rgb = None
            fills.append(rgb)
        opaque = all(f and len(f) == 8 and f.startswith("FF") for f in fills)
        if n == 0:
            self.add("R16", "%s 差异化柔和背景条件格式" % C.T2_GROUP_FIELD,
                     len(rules) == 0,
                     "源「%s」全空 → 分类数=0，规则数应为 0（实际 %d）" % (C.T2_GROUP_FIELD, len(rules)))
        else:
            self.add("R16", "%s 差异化柔和背景条件格式（%d 类 %d 条规则）" % (C.T2_GROUP_FIELD, n, len(rules)),
                     len(rules) >= n and opaque,
                     "规则数=%d 分类数=%d 填充=%s（须为 8 位 FF 开头的不透明 ARGB）"
                     % (len(rules), n, fills[:4]))

    # ---------- R17 条件格式公式锚点
    def r17_cf_formula(self):
        """
        条件格式公式必须锚定 sqref 左上角行：`EXACT($B3,"分类名")`。
        Excel 以区间左上角为基准做**相对行偏移**——若第 k 条规则写成自己那一行
        （`$B4`、`$B5`…），判断第 4 行时它会被偏移成 `$B5`，于是除首行外
        **一条规则都不命中**，分类底色露出斑马纹（2026-10-03 修复）。
        """
        ws = self.wbf[C.SHEET_REASON]
        if not C.REASON_SOFT_CF:
            n_any = sum(len(rng.rules) for rng in ws.conditional_formatting)
            self.add("R17", "%s 条件格式公式锚定区间首行：配置已关闭（无规则）" % C.T2_GROUP_FIELD,
                     n_any == 0, "规则数应为 0（实际 %d）" % n_any)
            return
        n_rules, bad = 0, []
        for rng in ws.conditional_formatting:
            base = min(cg.min_row for cg in rng.sqref.ranges)
            for rule in rng.rules:
                n_rules += 1
                for f in (rule.formula or []):
                    m = CF_FORMULA_RE.match(str(f).strip())
                    if not m:
                        bad.append((str(f), '须为 EXACT($列行,"分类") 形式'))
                    elif int(m.group(2)) != base:
                        bad.append((str(f), "锚点行 %s ≠ 区间首行 %d" % (m.group(2), base)))
        n = len(self.srcf["reasons"])
        if n == 0:
            self.add("R17", "%s 条件格式公式锚定区间首行" % C.T2_GROUP_FIELD, n_rules == 0,
                     "分类数=0 → 规则数应为 0（实际 %d）" % n_rules)
        else:
            self.add("R17", "%s 条件格式公式锚定区间首行（EXACT($B3,…)）" % C.T2_GROUP_FIELD,
                     not bad and n_rules >= n,
                     "规则数=%d 分类数=%d 不合规 %d 处 %s" % (n_rules, n, len(bad), bad[:3]))

    # ---------- R18 计数条件不得为裸区间
    def r18_criteria_guard(self):
        """
        COUNTIF/COUNTIFS 的「条件」参数不得是**裸区间**。
        裸区间的元素若是空白单元格，Excel 把该条件按 **0** 处理（而非空白），
        分母可能为 0 → `0/0` → `#DIV/0!`。「未回款原因分类」有 431 行空白，
        分组去重计数正踩此坑；而注入的缓存值是对的，只读缓存值的断言发现不了，
        故必须对公式文本直接断言（2026-10-03 新增）。
        """
        bad, checked = [], 0
        for name in self.profile["sheet_order"]:
            if name == C.SHEET_RAW:
                continue
            ws = self.wbf[name]
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if not (isinstance(v, str) and v.startswith("=")):
                        continue
                    for fn_name, args in iter_func_calls(v, ("COUNTIF", "COUNTIFS")):
                        if len(args) % 2:
                            continue
                        for idx in range(1, len(args), 2):
                            checked += 1
                            arg = args[idx].strip()
                            if BARE_RANGE_RE.match(arg):
                                bad.append((name, cell.coordinate, fn_name, arg))
        self.add("R18", "COUNTIF/COUNTIFS 条件参数均非裸区间（空白会被当作 0 → #DIV/0!）",
                 not bad, "检查 %d 处条件，违规 %d 处 %s" % (checked, len(bad), bad[:4]))

    def run_all(self):
        for fn in (self.r1_sheets, self.r2_raw_intact, self.r3_formulas, self.r4_errors,
                   self.r5_rowcounts, self.r6_values, self.r7_order, self.r8_formats,
                   self.r9_align, self.r10_header, self.r11_zebra, self.r12_border,
                   self.r13_freeze, self.r14_width, self.r15_title, self.r16_cf,
                   self.r17_cf_formula, self.r18_criteria_guard):
            fn()
        return self.results


def main():
    p = argparse.ArgumentParser(description="结果文件交付前自检")
    p.add_argument("--file", required=True, help="Skill 产出的结果 xlsx")
    p.add_argument("--source", required=True, help="原始输入 xlsx")
    p.add_argument("--config", default=DEFAULT_CONFIG_PATH,
                   help="业务配置 JSON（默认 config/contract_repayment.json）")
    p.add_argument("--profile", default=None,
                   help="期望值口径；不传则用配置里的 default_profile")
    p.add_argument("--json", default=None, help="把自检报告写成 JSON")
    args = p.parse_args()

    global C
    C = _Config(args.config)

    profile = C.get_profile(args.profile)
    ck = Checker(args.file, args.source, profile)
    try:
        results = ck.run_all()
    finally:
        ck.close()      # 无论通过与否都释放工作簿句柄
    passed = sum(1 for r in results if r["ok"])
    report = {
        "file": os.path.abspath(args.file),
        "source": os.path.abspath(args.source),
        "profile": profile["name"],
        "total_checks": len(results),
        "passed": passed,
        "failed": len(results) - passed,
        "all_passed": passed == len(results),
        "items": results,
    }
    text = json.dumps(report, ensure_ascii=False, indent=2, default=str)
    if args.json:
        with open(args.json, "w", encoding="utf-8") as fh:
            fh.write(text)
    for r in results:
        print("%s %s %s" % ("PASS" if r["ok"] else "FAIL", r["id"], r["desc"]))
        if not r["ok"]:
            print("      %s" % r["detail"])
    print("\n合计 %d 项，通过 %d 项，未通过 %d 项 → %s"
          % (len(results), passed, len(results) - passed,
             "全部合规" if passed == len(results) else "存在不合规项"))
    sys.exit(0 if passed == len(results) else 1)


if __name__ == "__main__":
    main()
